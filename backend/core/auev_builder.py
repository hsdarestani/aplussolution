import io
import re
from datetime import datetime, time, timedelta
from decimal import Decimal

from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Max
from django.utils import timezone
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .document_engine import convert_docx_to_pdf
from .models import AuevExport, AuevSetting, ClientCompany, Shift
from .shift_slots import ShiftSlot


DEFAULTS = {
    'permit_date': None,
    'framework_date': None,
    'effective_date': None,
    'required_qualification': '',
    'intended_activity': 'Servicekraft',
    'client_contract_text': '',
    'file_label': '',
}


def _iso(value):
    return value.isoformat() if value else ''


def _de(value):
    return value.strftime('%d.%m.%Y') if value else ''


def _date_range(start, end):
    tz = timezone.get_current_timezone()
    start_dt = timezone.make_aware(datetime.combine(start, time.min), tz)
    end_dt = timezone.make_aware(datetime.combine(end + timedelta(days=1), time.min), tz)
    return start_dt, end_dt


def _normalise_filename(value):
    value = re.sub(r'[\\/:*?"<>|]+', ' ', str(value or '')).strip()
    value = re.sub(r'\s+', ' ', value)
    return value[:120] or 'Kunde'


def _derive_file_label(client):
    name = (client.name or '').strip()
    match = re.search(r'\(([^()]+)\)', name)
    if match and len(match.group(1).strip()) <= 80:
        return match.group(1).strip()
    name = re.sub(r'\s+Catering\s*$', '', name, flags=re.I).strip()
    return name


def _default_setting():
    row = AuevSetting.objects.filter(client__isnull=True).order_by('created_at').first()
    if row:
        return row
    return AuevSetting.objects.create(
        intended_activity='Servicekraft',
    )


def resolve_settings(client):
    default = _default_setting()
    override = AuevSetting.objects.filter(client=client).first()

    def pick(field):
        value = getattr(override, field, None) if override else None
        if value not in (None, '', 0):
            return value
        value = getattr(default, field, None)
        if value not in (None, '', 0):
            return value
        return DEFAULTS[field]

    contract_text = pick('client_contract_text')
    if not contract_text:
        address = (client.address or '').strip()
        contract_text = client.name.strip()
        if address:
            contract_text = f'{contract_text}, {address}'

    file_label = pick('file_label') or _derive_file_label(client)
    return {
        'permit_date': pick('permit_date'),
        'framework_date': pick('framework_date'),
        'effective_date': pick('effective_date'),
        'required_qualification': str(pick('required_qualification') or ''),
        'intended_activity': str(pick('intended_activity') or 'Servicekraft'),
        'client_contract_text': str(contract_text).strip(),
        'file_label': _normalise_filename(file_label),
        'last_sequence_number': int(getattr(override, 'last_sequence_number', 0) or 0),
    }


def _workers_for_shift(shift):
    slots = list(
        shift.slots.filter(status=ShiftSlot.Status.CLAIMED, worker__isnull=False)
        .select_related('worker__user')
        .order_by('created_at')
    )
    if slots:
        return [slot.worker for slot in slots]
    return [shift.worker] if shift.worker_id else []


def _birth_date(worker):
    master = getattr(worker, 'master_data', None)
    raw = (master.data or {}).get('birth_date') if master else None
    if not raw:
        return ''
    if hasattr(raw, 'strftime'):
        return raw.strftime('%d.%m.%Y')
    text = str(raw).strip()
    for fmt in ('%d.%m.%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(text[:10], fmt).strftime('%d.%m.%Y')
        except ValueError:
            continue
    return text


def collect_rows(client, start, end):
    start_dt, end_dt = _date_range(start, end)
    shifts = list(
        Shift.objects.filter(
            client=client,
            starts_at__gte=start_dt,
            starts_at__lt=end_dt,
        )
        .exclude(status=Shift.Status.CANCELLED)
        .select_related('position', 'worker__user')
        .prefetch_related('slots__worker__user')
        .order_by('starts_at', 'ends_at', 'id')
    )

    rows = []
    unassigned = []
    missing_birth_dates = []
    shift_ids = []
    actual_dates = []

    for shift in shifts:
        shift_ids.append(str(shift.id))
        local_start = timezone.localtime(shift.starts_at)
        local_end = timezone.localtime(shift.ends_at)
        actual_dates.append(local_start.date())
        workers = _workers_for_shift(shift)
        missing_count = max(int(shift.required_count or 1) - len(workers), 0)
        if not workers or missing_count:
            unassigned.append({
                'shift_id': str(shift.id),
                'date': local_start.date().isoformat(),
                'start': local_start.strftime('%H:%M'),
                'position': shift.position.name,
                'missing_count': max(missing_count, 1 if not workers else 0),
            })

        for worker in workers:
            birth = _birth_date(worker)
            if not birth:
                missing_birth_dates.append(worker.user.get_full_name() or worker.user.email)
            last_name = (worker.user.last_name or '').strip()
            first_name = (worker.user.first_name or '').strip()
            if last_name or first_name:
                person = ', '.join(value for value in (last_name, first_name) if value)
            else:
                person = worker.user.get_full_name() or worker.user.email
            rows.append({
                'worker_id': str(worker.id),
                'name': person,
                'birth_date': birth,
                'name_birth': f'{person}, {birth or "–"}',
                'start': local_start.strftime('%H:%M'),
                'end': local_end.strftime('%H:%M'),
                'date': local_start.date().strftime('%d.%m.%Y'),
                'date_iso': local_start.date().isoformat(),
                'activity': shift.position.name,
                'shift_id': str(shift.id),
            })

    return {
        'shifts': shifts,
        'rows': rows,
        'shift_ids': shift_ids,
        'actual_dates': actual_dates,
        'unassigned': unassigned,
        'missing_birth_dates': sorted(set(missing_birth_dates)),
    }


def next_sequence_number(client):
    setting = AuevSetting.objects.filter(client=client).first()
    stored = int(setting.last_sequence_number if setting else 0)
    exported = AuevExport.objects.filter(client=client).aggregate(value=Max('sequence_number'))['value'] or 0
    return max(stored, int(exported)) + 1


def preview(client, start, end):
    collected = collect_rows(client, start, end)
    rows = collected['rows']
    dates = collected['actual_dates']
    settings = resolve_settings(client)
    sequence = next_sequence_number(client)
    weeks = sorted({date.isocalendar().week for date in dates})
    first = min(dates) if dates else None
    last = max(dates) if dates else None
    signature = first - timedelta(days=2) if first else None
    week_part = '-'.join(str(item) for item in weeks)
    file_stem = f'ANÜ - {settings["file_label"]} ({sequence}) KW{week_part}' if weeks else ''
    return {
        'client_id': str(client.id),
        'client_name': client.name,
        'date_from': start.isoformat(),
        'date_to': end.isoformat(),
        'first_shift_date': _iso(first),
        'last_shift_date': _iso(last),
        'signature_date_default': _iso(signature),
        'calendar_weeks': weeks,
        'sequence_number': sequence,
        'file_stem': file_stem,
        'shift_count': len(collected['shifts']),
        'row_count': len(rows),
        'unassigned': collected['unassigned'],
        'missing_birth_dates': collected['missing_birth_dates'],
        'new_template_limit_exceeded': len(rows) > 25,
        'settings': {
            'permit_date': _iso(settings['permit_date']),
            'framework_date': _iso(settings['framework_date']),
            'effective_date': _iso(settings['effective_date']),
            'required_qualification': settings['required_qualification'],
            'intended_activity': settings['intended_activity'],
            'client_contract_text': settings['client_contract_text'],
            'file_label': settings['file_label'],
        },
    }


def _font(run, size=10.5, bold=False, italic=False, underline=False):
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    return run


def _paragraph(container, text='', size=10.5, bold=False, before=0, after=0, keep=False):
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = keep
    _font(paragraph.add_run(text), size=size, bold=bold)
    return paragraph


def _set_cell_text(cell, text, size=9.0, bold=False, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    _font(p.add_run(str(text or '')), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge not in kwargs:
            continue
        edge_data = kwargs[edge]
        tag = 'w:{}'.format(edge)
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key, value in edge_data.items():
            element.set(qn('w:{}'.format(key)), str(value))


def _no_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(
                cell,
                top={'val': 'nil'},
                bottom={'val': 'nil'},
                left={'val': 'nil'},
                right={'val': 'nil'},
            )


def _bottom_border(cell):
    _set_cell_border(
        cell,
        top={'val': 'nil'},
        left={'val': 'nil'},
        right={'val': 'nil'},
        bottom={'val': 'single', 'sz': '4', 'color': '777777'},
    )


def _page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def _base_document(landscape=False):
    doc = Document()
    section = doc.sections[0]
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(10.5)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(0.79)
        section.bottom_margin = Cm(0.79)
        section.left_margin = Cm(0.79)
        section.right_margin = Cm(1.6)
    else:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
    return doc


def _classic_docx(data):
    doc = _base_document(False)
    section = doc.sections[0]
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(header_p.add_run(str(data['sequence_number'])), size=18, bold=True)
    footer_p = section.footer.paragraphs[0]
    _page_field(footer_p)

    _paragraph(doc, 'Anhang 1: Einzelarbeitnehmerüberlassungsvertrag', size=13, bold=True, after=3)
    _paragraph(doc, 'Zwischen', after=1)

    client_table = doc.add_table(rows=1, cols=1)
    client_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _set_cell_text(client_table.cell(0, 0), data['settings']['client_contract_text'], size=10.5)
    _bottom_border(client_table.cell(0, 0))

    _paragraph(doc, 'und', before=1, after=1)
    _paragraph(doc, 'A+ Solution GmbH, Carl-Sonnenschein Straße 57, 65936 Frankfurt am Main', after=0)
    _paragraph(doc, '(Personaldienstleister) wird folgender Arbeitnehmerüberlassungsvertrag geschlossen:', after=5)

    _paragraph(doc, '§ 1 Erlaubnis zur Arbeitnehmerüberlassung', bold=True, keep=True)
    permit = data['settings']['permit_date_de'] or '__________'
    _paragraph(
        doc,
        'Der Personaldienstleister erklärt, im Besitz einer befristeten Erlaubnis zur Arbeitnehmerüberlassung zu sein, '
        'zuletzt erteilt und nicht widerrufen von der Bundesagentur für Arbeit, Agentur für Arbeit Düsseldorf am '
        f'{permit} in Düsseldorf.',
        after=4,
    )

    _paragraph(doc, '§ 2 Rahmenvereinbarung', bold=True, keep=True)
    framework = data['settings']['framework_date_de'] or '__________'
    _paragraph(
        doc,
        'Die Rahmenvereinbarung zur Arbeitnehmerüberlassung vom '
        f'{framework} zwischen Auftraggeber und Personaldienstleister findet auf diesen Arbeitnehmerüberlassungsvertrag Anwendung.',
        after=4,
    )

    _paragraph(doc, '§ 3 Gegenstand des Vertrages / Überlassungsbedingungen', bold=True, keep=True)
    effective = data['settings']['effective_date_de'] or data['first_shift_date_de']
    _paragraph(
        doc,
        'Der Personaldienstleister überlässt mit Wirkung zum '
        f'{effective} an den Auftraggeber folgende Zeitarbeitnehmer an den in § 2 Absatz 2 der Rahmenvereinbarung festgelegten Betrieb.',
        after=4,
    )

    meta = doc.add_table(rows=2, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.columns[0].width = Cm(6.7)
    meta.columns[1].width = Cm(10.1)
    _no_table_borders(meta)
    _set_cell_text(meta.cell(0, 0), 'Vorgesehene Tätigkeit:', size=10.5)
    _bottom_border(meta.cell(0, 0))
    _set_cell_text(meta.cell(0, 1), data['settings']['intended_activity'], size=10.5)
    _set_cell_text(meta.cell(1, 0), 'Betriebliche Arbeitszeit in Stunden/MA:', size=10.5)
    _bottom_border(meta.cell(1, 0))
    _set_cell_text(meta.cell(1, 1), '(S.u.)', size=10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [7.1, 1.45, 1.45, 2.2, 3.03, 1.57]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Cm(width)
    headers = ['Name, Vorname, Geburtsdatum', 'Start', 'Ende', 'Datum', 'Tätigkeit', '']
    for idx, value in enumerate(headers):
        _set_cell_text(table.cell(0, idx), value, size=9.0, bold=True)
    for row in data['rows']:
        cells = table.add_row().cells
        values = [row['name_birth'], row['start'], row['end'], row['date'], row['activity'], '']
        for idx, value in enumerate(values):
            _set_cell_text(cells[idx], value, size=8.8)
    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    _paragraph(
        doc,
        '(   ) Konkretisierung zum aktuellen Zeitpunkt nicht bekannt. Wird rechtzeitig per E-Mail mitgeteilt.',
        size=10,
        before=4,
        after=5,
    )

    clauses = [
        ('(1)', 'Die namentliche Nennung und die Angabe des Geburtsdatums erfolgt ausschließlich hinsichtlich § 1 Abs. 1 Satz 6 AÜG. '
                'Sollte die Person des Zeitarbeitnehmers im Zeitpunkt des Abschlusses des Einzelarbeitnehmerüberlassungsvertrages noch unbekannt sein, '
                'so ist der Zeitarbeitnehmer von Auftraggeber und Personaldienstleister rechtzeitig vor Einsatzbeginn namentlich unter Angabe des Geburtsdatums '
                'einvernehmlich zu benennen (Konkretisierung).'),
        ('(2)', 'Die Überlassungsvergütung richtet sich nach der tatsächlichen Arbeitszeit der eingesetzten Arbeitnehmer, mindestens aber nach der in Absatz 1 genannten betrieblichen Arbeitszeit.'),
        ('(4)', 'Es werden folgende Zuschläge vereinbart:'),
    ]
    for number, text in clauses:
        t = doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        _no_table_borders(t)
        t.columns[0].width = Cm(0.8)
        t.columns[1].width = Cm(16.0)
        _set_cell_text(t.cell(0, 0), number, size=10)
        _set_cell_text(t.cell(0, 1), text, size=10)

    _paragraph(doc, '§ 4 Arbeitsschutz', bold=True, before=5, keep=True)
    _paragraph(doc, '(1) Bitte Zutreffendes ankreuzen:', after=1)
    _paragraph(doc, '□  Für den Einsatz der überlassenen Zeitarbeitnehmer sind keine arbeitsmedizinischen Vorsorgeuntersuchungen erforderlich. (X)', after=1)
    _paragraph(doc, '□  Für den Einsatz der überlassenen Zeitarbeitnehmer sind folgende arbeitsmedizinischen Vorsorgeuntersuchungen erforderlich:', after=1)
    med = doc.add_table(rows=1, cols=2)
    _no_table_borders(med)
    med.columns[0].width = Cm(3.0)
    med.columns[1].width = Cm(13.8)
    _set_cell_text(med.cell(0, 0), 'Angabe:', size=10.5)
    _set_cell_text(med.cell(0, 1), 'hier eintragen', size=10.5)
    _bottom_border(med.cell(0, 1))
    _paragraph(doc, 'Diese werden vom Personaldienstleister vor Überlassungsbeginn durchgeführt und dem Auftraggeber nachgewiesen.', after=4)

    _paragraph(doc, '§ 5 Befristung', bold=True, keep=True)
    end_table = doc.add_table(rows=1, cols=2)
    _no_table_borders(end_table)
    end_table.columns[0].width = Cm(13.0)
    end_table.columns[1].width = Cm(3.8)
    _set_cell_text(end_table.cell(0, 0), 'Dieser Einzelarbeitnehmerüberlassungsvertrag wird zunächst befristet bis zum', size=10.5)
    _set_cell_text(end_table.cell(0, 1), data['last_shift_date_de'], size=10.5)
    _bottom_border(end_table.cell(0, 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    sig = doc.add_table(rows=2, cols=3)
    _no_table_borders(sig)
    sig.columns[0].width = Cm(7.8)
    sig.columns[1].width = Cm(1.2)
    sig.columns[2].width = Cm(7.8)
    _set_cell_text(sig.cell(0, 0), '', size=10.5)
    _bottom_border(sig.cell(0, 0))
    _set_cell_text(sig.cell(0, 2), data['signature_date_de'], size=10.5)
    _bottom_border(sig.cell(0, 2))
    _set_cell_text(sig.cell(1, 0), '[Datum, Unterschrift Auftraggeber]', size=9.5)
    _set_cell_text(sig.cell(1, 2), '[Datum, Unterschrift Personaldienstleister]', size=9.5)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _add_cell_paragraph(cell, text='', size=11, bold=False, before=0, after=0):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    _font(p.add_run(text), size=size, bold=bold)
    return p


def _new_docx(data):
    if len(data['rows']) > 25:
        raise ValueError('Die neue Vorlage bietet Platz für maximal 25 Mitarbeiterzeilen.')

    doc = _base_document(True)
    section = doc.sections[0]
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(footer_p.add_run(str(data['sequence_number'])), size=18, bold=True)

    outer = doc.add_table(rows=1, cols=2)
    outer.alignment = WD_TABLE_ALIGNMENT.CENTER
    outer.autofit = False
    outer.columns[0].width = Cm(12.2)
    outer.columns[1].width = Cm(13.9)
    _no_table_borders(outer)
    left = outer.cell(0, 0)
    right = outer.cell(0, 1)
    left.text = ''
    right.text = ''

    _add_cell_paragraph(left, 'Anhang 1: Einzelarbeitnehmerüberlassungsvertrag', size=13, bold=True, after=5)
    _add_cell_paragraph(left, 'Zwischen', size=11, after=1)
    client_p = _add_cell_paragraph(left, data['settings']['client_contract_text'], size=10.5, after=1)
    pPr = client_p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:color'), '777777')
    pBdr.append(bottom)
    pPr.append(pBdr)
    _add_cell_paragraph(left, 'und', size=11, after=1)
    _add_cell_paragraph(left, 'A+ Solution GmbH, Carl-Sonnenschein Straße 57, 65936 Frankfurt a.M.', size=10.5, after=0)
    _add_cell_paragraph(left, '(Personaldienstleister)', size=10.5, after=0)
    _add_cell_paragraph(left, 'wird folgender Arbeitnehmerüberlassungsvertrag geschlossen:', size=10.5, after=8)

    _add_cell_paragraph(left, '§ 1 Erlaubnis zur Arbeitnehmerüberlassung', size=10.5, bold=True, after=1)
    permit = data['settings']['permit_date_de'] or '__________'
    _add_cell_paragraph(
        left,
        'Der Personaldienstleister erklärt, im Besitz einer befristeten Erlaubnis zur Arbeitnehmerüberlassung zu sein, '
        'zuletzt erteilt und nicht widerrufen von der Bundesagentur für Arbeit, Agentur für Arbeit Düsseldorf am '
        f'{permit} in Düsseldorf.',
        size=10.5,
        after=7,
    )

    _add_cell_paragraph(left, '§ 2 Rahmenvereinbarung', size=10.5, bold=True, after=1)
    framework = data['settings']['framework_date_de'] or '__________'
    _add_cell_paragraph(
        left,
        'Die Rahmenvereinbarung zur Arbeitnehmerüberlassung vom '
        f'{framework} zwischen Auftraggeber und Personaldienstleister findet auf diesen Arbeitnehmerüberlassungsvertrag Anwendung.',
        size=10.5,
        after=7,
    )

    _add_cell_paragraph(left, '§ 3 Gegenstand des Vertrages / Überlassungsbedingungen', size=10.5, bold=True, after=1)
    effective = data['settings']['effective_date_de'] or data['first_shift_date_de']
    _add_cell_paragraph(
        left,
        'Der Personaldienstleister überlässt mit Wirkung zum '
        f'{effective} an den Auftraggeber folgende Zeitarbeitnehmer an den in § 2 Absatz 2 der Rahmenvereinbarung festgelegten Betrieb.',
        size=10.5,
        after=18,
    )

    _add_cell_paragraph(
        left,
        f'Dieser Einzelarbeitnehmerüberlassungsvertrag wird zunächst befristet bis zum  {data["last_shift_date_de"]}',
        size=10.5,
        after=12,
    )
    _add_cell_paragraph(left, data['signature_date_de'], size=10.5, after=22)

    sig = left.add_table(rows=2, cols=2)
    _no_table_borders(sig)
    _set_cell_text(sig.cell(0, 0), '', size=9)
    _set_cell_text(sig.cell(0, 1), '', size=9)
    _bottom_border(sig.cell(0, 0))
    _bottom_border(sig.cell(0, 1))
    _set_cell_text(sig.cell(1, 0), '[Datum, Unterschrift Personaldienstleister]', size=8.8)
    _set_cell_text(sig.cell(1, 1), '[Datum, Unterschrift Auftraggeber]', size=8.8)

    employee = right.add_table(rows=26, cols=6)
    employee.alignment = WD_TABLE_ALIGNMENT.RIGHT
    employee.autofit = False
    widths = [0.78, 6.43, 1.37, 1.37, 2.08, 2.57]
    for idx, width in enumerate(widths):
        employee.columns[idx].width = Cm(width)
    headers = ['', 'Name, Vorname, Geburtsdatum', 'Start', 'Ende', 'Datum', 'Tätigkeit']
    for idx, value in enumerate(headers):
        _set_cell_text(employee.cell(0, idx), value, size=8.2, bold=True)
    for index in range(25):
        cells = employee.rows[index + 1].cells
        if index < len(data['rows']):
            row = data['rows'][index]
            values = [str(index + 1), row['name_birth'], row['start'], row['end'], row['date'], row['activity']]
        else:
            values = ['', '', '', '', '', '']
        for col, value in enumerate(values):
            _set_cell_text(cells[col], value, size=8.0)
        employee.rows[index + 1].height = Cm(0.67)
        employee.rows[index + 1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _document_data(client, start, end, signature_date, sequence_number):
    collected = collect_rows(client, start, end)
    if not collected['shifts']:
        raise ValueError('Im gewählten Zeitraum wurden für diesen Kunden keine Einsätze gefunden.')
    if collected['unassigned']:
        raise ValueError('Mindestens ein Einsatz ist noch nicht vollständig mit Mitarbeitern besetzt.')
    if collected['missing_birth_dates']:
        names = ', '.join(collected['missing_birth_dates'][:8])
        raise ValueError(f'Geburtsdatum fehlt für: {names}')

    settings = resolve_settings(client)
    dates = collected['actual_dates']
    first = min(dates)
    last = max(dates)
    weeks = sorted({value.isocalendar().week for value in dates})
    if signature_date > first:
        raise ValueError('Das Unterschriftsdatum darf nicht nach dem ersten Einsatz liegen.')

    return {
        'client': client,
        'rows': collected['rows'],
        'shift_ids': collected['shift_ids'],
        'settings': {
            **settings,
            'permit_date_de': _de(settings['permit_date']),
            'framework_date_de': _de(settings['framework_date']),
            'effective_date_de': _de(settings['effective_date']),
        },
        'first_shift_date': first,
        'last_shift_date': last,
        'first_shift_date_de': _de(first),
        'last_shift_date_de': _de(last),
        'signature_date': signature_date,
        'signature_date_de': _de(signature_date),
        'weeks': weeks,
        'sequence_number': sequence_number,
    }


@transaction.atomic
def generate_export(*, client, start, end, template_key, signature_date, actor=None, sequence_number=None):
    if template_key not in {AuevExport.Template.CLASSIC, AuevExport.Template.NEW}:
        raise ValueError('Unbekannte ANÜ Vorlage.')

    setting, _ = AuevSetting.objects.select_for_update().get_or_create(client=client)
    max_export = AuevExport.objects.filter(client=client).aggregate(value=Max('sequence_number'))['value'] or 0
    current = max(int(setting.last_sequence_number or 0), int(max_export))
    sequence = int(sequence_number or current + 1)
    if sequence < 1:
        raise ValueError('Die Dokumentnummer muss mindestens 1 sein.')
    if AuevExport.objects.filter(client=client, sequence_number=sequence).exists():
        raise ValueError(f'Die ANÜ Nummer {sequence} existiert für diesen Kunden bereits.')

    data = _document_data(client, start, end, signature_date, sequence)
    if template_key == AuevExport.Template.NEW:
        docx_bytes = _new_docx(data)
    else:
        docx_bytes = _classic_docx(data)
    pdf_bytes = convert_docx_to_pdf(docx_bytes)

    weeks = data['weeks']
    stem = f'ANÜ - {data["settings"]["file_label"]} ({sequence}) KW{"-".join(str(item) for item in weeks)}'
    stem = _normalise_filename(stem)

    export = AuevExport(
        client=client,
        template_key=template_key,
        date_from=start,
        date_to=end,
        first_shift_date=data['first_shift_date'],
        last_shift_date=data['last_shift_date'],
        signature_date=signature_date,
        sequence_number=sequence,
        calendar_weeks=weeks,
        settings_snapshot={
            'permit_date': _iso(data['settings']['permit_date']),
            'framework_date': _iso(data['settings']['framework_date']),
            'effective_date': _iso(data['settings']['effective_date']),
            'required_qualification': data['settings']['required_qualification'],
            'intended_activity': data['settings']['intended_activity'],
            'client_contract_text': data['settings']['client_contract_text'],
            'file_label': data['settings']['file_label'],
        },
        shift_ids=data['shift_ids'],
        row_count=len(data['rows']),
        file_stem=stem,
        created_by=actor,
    )
    export.docx.save(f'{stem}.docx', ContentFile(docx_bytes), save=False)
    export.pdf.save(f'{stem}.pdf', ContentFile(pdf_bytes), save=False)
    export.save()

    if sequence > setting.last_sequence_number:
        setting.last_sequence_number = sequence
        setting.save(update_fields=['last_sequence_number', 'updated_at'])

    return export


def export_dict(item):
    return {
        'id': str(item.id),
        'client_id': str(item.client_id),
        'client_name': item.client.name,
        'template_key': item.template_key,
        'template_label': item.get_template_key_display(),
        'date_from': item.date_from,
        'date_to': item.date_to,
        'first_shift_date': item.first_shift_date,
        'last_shift_date': item.last_shift_date,
        'signature_date': item.signature_date,
        'sequence_number': item.sequence_number,
        'calendar_weeks': item.calendar_weeks,
        'row_count': item.row_count,
        'file_stem': item.file_stem,
        'docx_url': item.docx.url if item.docx else '',
        'pdf_url': item.pdf.url if item.pdf else '',
        'created_at': item.created_at,
    }
