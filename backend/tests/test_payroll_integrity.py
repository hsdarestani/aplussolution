from datetime import date, datetime, time, timedelta
from decimal import Decimal

import pytest
from django.utils import timezone

from core.models import Document, EmployeeMasterData, PayrollStatement, Shift, TimeEntry, TimeOffRequest, WorkingTimeAccountRecord, WorkingTimeSetting
from core.native_cutover import sync_working_time
from core.payroll_engine import effective_hourly_rate


@pytest.mark.django_db
def test_fixed_salary_balance_uses_contractual_soll_not_paid_hours(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('100.00')
    worker.tariff_hourly_rate = Decimal('0.00')
    worker.save(update_fields=['monthly_hours', 'tariff_hourly_rate', 'updated_at'])
    EmployeeMasterData.objects.create(
        worker=worker,
        data={'compensation_type': 'salary', 'monthly_salary': '2975.00'},
    )
    WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('100.00'),
        hourly_rate=Decimal('0.00'),
    )

    start = timezone.make_aware(
        datetime(2026, 9, 1, 8, 0),
        timezone.get_current_timezone(),
    )
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        approved=True,
    )

    sync_working_time(date(2026, 9, 1), date(2026, 9, 30))
    record = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=date(2026, 9, 1))

    assert record.ist_hours == Decimal('8.00')
    assert record.soll_hours == Decimal('100.00')
    assert record.paid_total_hours == Decimal('0.00')
    assert record.saldo_cumulative == Decimal('-92.00')
    assert record.gross_amount == Decimal('2975.00')

    from core.working_time import record_dict
    data = record_dict(record)
    assert data['balance_basis'] == 'soll_salary'
    assert data['balance_reference_hours'] == '100.00'
    assert data['monthly_balance_hours'] == '-92.00'


@pytest.mark.django_db
def test_payroll_excludes_unapproved_time_and_includes_allowance(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('10.00')
    worker.tariff_hourly_rate = Decimal('15.50')
    worker.extra_allowance = Decimal('2.00')
    worker.save(update_fields=[
        'monthly_hours', 'tariff_hourly_rate', 'extra_allowance', 'updated_at'
    ])

    today = timezone.localdate()
    start = timezone.make_aware(
        datetime.combine(today, time(8, 0)), timezone.get_current_timezone()
    )
    approved_shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    unapproved_shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start + timedelta(hours=10),
        ends_at=start + timedelta(hours=14),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=approved_shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        approved=True,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=unapproved_shift,
        clock_in=start + timedelta(hours=10),
        clock_out=start + timedelta(hours=14),
        approved=False,
    )

    log = sync_working_time(today, today)
    record = WorkingTimeAccountRecord.objects.get(
        worker=worker,
        year_month=today.replace(day=1),
    )

    assert record.ist_hours == Decimal('8.00')
    assert record.soll_hours == Decimal('10.00')
    assert record.difference_hours == Decimal('-2.00')
    assert record.hourly_rate == Decimal('17.50')
    assert record.gross_amount == Decimal('140.00')
    assert len(record.raw_entries) == 1
    assert record.raw_entries[0]['approved'] is True
    assert record.source == 'aplus_time_entries'

    assert log.status == 'warning'
    assert log.metadata['source'] == 'aplus_time_entries'
    assert log.metadata['closed_entries'] == 2
    assert log.metadata['approved_entries'] == 1
    assert log.metadata['excluded_unapproved_entries'] == 1
    assert 'aus der Lohnvorbereitung ausgeschlossen' in log.message


@pytest.mark.django_db
def test_payroll_includes_entry_after_manager_approval(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('8.00')
    worker.tariff_hourly_rate = Decimal('20.00')
    worker.save(update_fields=['monthly_hours', 'tariff_hourly_rate', 'updated_at'])

    today = timezone.localdate()
    start = timezone.make_aware(
        datetime.combine(today, time(9, 0)), timezone.get_current_timezone()
    )
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=4),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    entry = TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=4),
        approved=False,
    )

    sync_working_time(today, today)
    record = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=today.replace(day=1))
    assert record.ist_hours == Decimal('0.00')
    assert record.gross_amount == Decimal('0.00')

    entry.approved = True
    entry.save(update_fields=['approved', 'updated_at'])
    sync_working_time(today, today)
    record.refresh_from_db()
    assert record.ist_hours == Decimal('4.00')
    assert record.gross_amount == Decimal('80.00')


@pytest.mark.django_db
def test_rebuild_all_creates_month_for_closed_pending_attendance(
    auth_admin, worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('8.00')
    worker.save(update_fields=['monthly_hours', 'updated_at'])

    today = timezone.localdate()
    start = timezone.make_aware(
        datetime.combine(today, time(9, 0)), timezone.get_current_timezone()
    )
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=4),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=4),
        approved=False,
    )

    response = auth_admin.post('/api/working-time/rebuild-all/', {}, format='json')
    assert response.status_code == 200
    assert response.data['records_count'] >= 1

    record = WorkingTimeAccountRecord.objects.get(
        worker=worker,
        year_month=today.replace(day=1),
    )
    assert record.ist_hours == Decimal('0.00')
    assert record.raw_entries == []

    records_response = auth_admin.get('/api/working-time/records/')
    assert records_response.status_code == 200
    assert records_response.data['count'] >= 1
    assert any(item['id'] == str(record.id) for item in records_response.data['results'])


@pytest.mark.django_db
def test_payroll_setting_is_base_rate_and_allowance_is_added(worker_user):
    worker = worker_user.worker_profile
    worker.tariff_hourly_rate = Decimal('15.00')
    worker.extra_allowance = Decimal('2.50')
    worker.save(update_fields=['tariff_hourly_rate', 'extra_allowance', 'updated_at'])
    setting = WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('80.00'),
        hourly_rate=Decimal('18.00'),
        active=True,
    )

    base, allowance, effective = effective_hourly_rate(worker, setting)
    assert base == Decimal('18.00')
    assert allowance == Decimal('2.50')
    assert effective == Decimal('20.50')



@pytest.mark.django_db
def test_fifty_worked_thirty_eight_paid_leaves_twelve_hours_credit(
    auth_admin, worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('38.00')
    worker.tariff_hourly_rate = Decimal('15.00')
    worker.save(update_fields=['monthly_hours', 'tariff_hourly_rate', 'updated_at'])

    month = timezone.localdate().replace(day=1)
    tz = timezone.get_current_timezone()
    for offset in range(5):
        actual_start = timezone.make_aware(
            datetime.combine(month + timedelta(days=offset), time(8, 0)), tz
        )
        shift = Shift.objects.create(
            client=company,
            location=location,
            position=position,
            worker=worker,
            # The plan is intentionally only eight hours; actual attendance is ten.
            starts_at=actual_start,
            ends_at=actual_start + timedelta(hours=8),
            break_minutes=0,
            status=Shift.Status.CONFIRMED,
        )
        TimeEntry.objects.create(
            worker=worker,
            shift=shift,
            clock_in=actual_start,
            clock_out=actual_start + timedelta(hours=10),
            break_minutes=0,
            approved=True,
        )

    sync_working_time(month, month + timedelta(days=6))
    record = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=month)

    assert record.ist_hours == Decimal('50.00')
    assert record.soll_hours == Decimal('38.00')
    assert record.paid_total_hours == Decimal('38.00')
    assert record.saldo_cumulative == Decimal('12.00')
    assert record.employment_type_snapshot == 'minijob'
    assert len(record.raw_entries) == 5
    assert all(item['client_name'] == company.name for item in record.raw_entries)
    assert all(item['worked_minutes'] == 600 for item in record.raw_entries)
    assert all(item['planned_end'] != item['clock_out'] for item in record.raw_entries)

    response = auth_admin.patch(
        f'/api/working-time/records/{record.id}/',
        {'paid_total_hours': '45.00'},
        format='json',
    )
    assert response.status_code == 200
    record.refresh_from_db()
    assert record.paid_total_hours == Decimal('45.00')
    assert record.saldo_cumulative == Decimal('5.00')


@pytest.mark.django_db
def test_payroll_repairs_implausible_one_day_clockout_rollover(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    day = date(2026, 9, 16)
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(day, time(16, 0)), tz)
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=6),
        break_minutes=30,
        status=Shift.Status.CONFIRMED,
    )
    entry = TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(days=1, hours=6, minutes=30),
        break_minutes=30,
        approved=True,
    )

    sync_working_time(day, day)
    record = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=day.replace(day=1))

    assert record.ist_hours == Decimal('6.00')
    assert len(record.raw_entries) == 1
    audit_row = record.raw_entries[0]
    assert audit_row['worked_minutes'] == 360
    assert audit_row['night_minutes'] == 0
    assert audit_row['clock_out_rollover_corrected'] is True
    source_clock_out = datetime.fromisoformat(audit_row['source_clock_out'])
    assert source_clock_out == entry.clock_out
    corrected_clock_out = datetime.fromisoformat(audit_row['local_clock_out'])
    assert corrected_clock_out.astimezone(tz).date() == day
    assert corrected_clock_out.astimezone(tz).time().replace(tzinfo=None) == time(22, 30)


@pytest.mark.django_db
def test_historical_wiw_time_is_authoritative_even_without_native_approval(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('8.00')
    worker.save(update_fields=['monthly_hours', 'updated_at'])

    day = timezone.localdate().replace(day=1)
    start = timezone.make_aware(
        datetime.combine(day, time(22, 0)), timezone.get_current_timezone()
    )
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        break_minutes=0,
        approved=False,
        wiw_time_id='historical-1',
    )

    sync_working_time(day, day + timedelta(days=1))
    record = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=day)

    assert record.ist_hours == Decimal('8.00')
    assert record.raw_entries[0]['source'] == 'wiw_historical'
    assert record.raw_entries[0]['client_name'] == company.name
    assert record.raw_entries[0]['night_minutes'] > 0


@pytest.mark.django_db
def test_one_time_lexware_bank_import_links_transfer_to_employee_month(
    auth_admin, worker_user
):
    from django.core.files.uploadedfile import SimpleUploadedFile

    period = timezone.localdate().replace(day=1)
    csv_body = (
        '"EXTF";700;21;"Buchungsstapel";13;;;;;;;1001;456;;;;;;;1;;0;"EUR"\n'
        'Umsatz (ohne Soll/Haben-Kz);Soll/Haben-Kennzeichen;Belegdatum;Buchungstext\n'
        f'500,00;S;05{period:%m};Gehalt Anna Becker {period:%m/%Y}\n'
    ).encode('utf-8')
    upload = SimpleUploadedFile('lexware.csv', csv_body, content_type='text/csv')

    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': period.strftime('%Y-%m'), 'file': upload},
        format='multipart',
    )
    assert response.status_code == 200
    assert response.data['employees'][0]['employee_name'] == 'Anna Becker'
    assert response.data['employees'][0]['transferred_amount'] == '500.00'

    statement = PayrollStatement.objects.get(
        worker=worker_user.worker_profile,
        period=period,
    )
    assert statement.transferred_amount == Decimal('500.00')
    assert statement.source == 'lexware_bank_export'
    assert statement.payment_date == period.replace(day=5)
    assert len(statement.raw_data) == 1



@pytest.mark.django_db
def test_lexware_import_rejects_filename_period_mismatch(auth_admin):
    from django.core.files.uploadedfile import SimpleUploadedFile

    upload = SimpleUploadedFile(
        'Lohnabrechnungen_2026-09.pdf',
        b'not parsed because period validation runs first',
        content_type='application/pdf',
    )
    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-10', 'file': upload},
        format='multipart',
    )
    assert response.status_code == 400
    assert response.data['detected_period'] == '2026-09'
    assert '2026-10' in response.data['detail']


@pytest.mark.django_db
def test_lexware_employee_master_data_import_updates_personal_and_payroll_settings(
    auth_admin, worker_user
):
    import json

    from django.core.files.uploadedfile import SimpleUploadedFile

    worker = worker_user.worker_profile
    upload = SimpleUploadedFile(
        'lexware-stammdaten.json',
        json.dumps({
            'employees': [{
                'name': worker.user.get_full_name(),
                'app_employment_type': 'minijob',
                'data': {
                    'employment_type_lexware': 'Minijobber – Rentenversicherungsfrei',
                    'entry_date': '2026-08-30',
                    'weekly_hours': '8.00',
                    'compensation_type': 'hourly',
                    'hourly_rate': '16.00',
                    'iban': 'DE00 TEST',
                    'night_surcharge_percent': '25',
                    'sunday_surcharge_percent': '50',
                },
            }],
        }).encode('utf-8'),
        content_type='application/json',
    )

    response = auth_admin.post(
        '/api/workers/master-data/import/',
        {'file': upload},
        format='multipart',
    )
    assert response.status_code == 200
    assert len(response.data['employees']) == 1
    assert response.data['unmatched'] == []

    worker.refresh_from_db()
    assert worker.employment_type == 'minijob'
    assert worker.tariff_hourly_rate == Decimal('16.00')
    assert worker.monthly_hours == Decimal('34.67')

    master = worker.master_data
    assert master.data['entry_date'] == '2026-08-30'
    assert master.data['iban'] == 'DE00 TEST'
    assert master.source_map['iban'] == 'lexware_stammdaten'

    setting = worker.working_time_setting
    assert setting.monthly_limit == Decimal('34.67')
    assert setting.hourly_rate == Decimal('16.00')
    assert setting.night_surcharge_percent == Decimal('25')
    assert setting.sunday_surcharge_percent == Decimal('50')


@pytest.mark.django_db
def test_lexware_master_data_import_can_match_by_self_service_email(
    auth_admin, worker_user
):
    import json

    from django.core.files.uploadedfile import SimpleUploadedFile

    worker = worker_user.worker_profile
    worker.user.email = 'francesco@example.com'
    worker.user.first_name = 'Legacy'
    worker.user.last_name = 'Name'
    worker.user.save(update_fields=['email', 'first_name', 'last_name'])

    upload = SimpleUploadedFile(
        'lexware-stammdaten.json',
        json.dumps({
            'employees': [{
                'name': 'Francesco Trulli',
                'app_employment_type': 'vollzeit',
                'data': {
                    'self_service_email': 'francesco@example.com',
                    'weekly_hours': '38.50',
                    'compensation_type': 'salary',
                    'monthly_salary': '2975.00',
                },
            }],
        }).encode('utf-8'),
        content_type='application/json',
    )

    response = auth_admin.post(
        '/api/workers/master-data/import/',
        {'file': upload},
        format='multipart',
    )
    assert response.status_code == 200
    assert len(response.data['employees']) == 1
    assert response.data['unmatched'] == []

    worker.refresh_from_db()
    assert worker.employment_type == 'vollzeit'
    assert worker.monthly_hours == Decimal('166.83')
    assert worker.master_data.data['compensation_type'] == 'salary'
    assert worker.master_data.data['monthly_salary'] == '2975.00'


@pytest.mark.django_db
def test_refresh_contract_terms_replaces_stale_closed_month_snapshot(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.employment_type = 'vollzeit'
    worker.monthly_hours = Decimal('166.83')
    worker.save(update_fields=['employment_type', 'monthly_hours', 'updated_at'])
    EmployeeMasterData.objects.create(
        worker=worker,
        data={'compensation_type': 'salary', 'monthly_salary': '2975.00'},
    )
    WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('166.83'),
        hourly_rate=Decimal('0.00'),
    )

    start = timezone.make_aware(
        datetime(2026, 9, 1, 8, 0),
        timezone.get_current_timezone(),
    )
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        approved=True,
    )
    stale = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('8.00'),
        soll_hours=Decimal('0.00'),
        paid_total_hours=Decimal('0.00'),
        employment_type_snapshot='minijob',
        saldo_cumulative=Decimal('8.00'),
    )

    sync_working_time(
        date(2026, 9, 1),
        date(2026, 9, 30),
        refresh_contract_terms=True,
    )
    stale.refresh_from_db()

    assert stale.soll_hours == Decimal('166.83')
    assert stale.employment_type_snapshot == 'vollzeit'
    assert stale.saldo_cumulative == Decimal('-158.83')
    assert stale.gross_amount == Decimal('2975.00')


@pytest.mark.django_db
def test_editable_worktime_docx_export(auth_admin, worker_user):
    import io
    import zipfile

    worker = worker_user.worker_profile
    period = date(2026, 9, 1)
    WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=period,
        ist_hours=Decimal('50.00'),
        soll_hours=Decimal('40.00'),
        paid_total_hours=Decimal('38.00'),
        saldo_cumulative=Decimal('12.00'),
        raw_entries=[{
            'local_clock_in': '2026-09-01T08:00:00+02:00',
            'local_clock_out': '2026-09-01T16:30:00+02:00',
            'planned_start': '2026-09-01T08:00:00+02:00',
            'planned_end': '2026-09-01T16:00:00+02:00',
            'break_minutes': 30,
            'worked_minutes': 480,
            'night_minutes': 0,
            'saturday_minutes': 0,
            'sunday_minutes': 0,
            'client_name': 'Testkunde',
        }],
    )

    response = auth_admin.get(f'/api/working-time/docx/{worker.id}/')
    assert response.status_code == 200
    assert response['Content-Type'].startswith(
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    assert '01_Arbeitszeitnachweis' in response['Content-Disposition']
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        document_xml = archive.read('word/document.xml').decode('utf-8')
    assert 'Arbeitszeitnachweis und Lohnkonto' in document_xml
    assert 'Testkunde' in document_xml


@pytest.mark.django_db
def test_lexware_reconciliation_docx_export(auth_admin, worker_user):
    import io
    import zipfile

    worker = worker_user.worker_profile
    period = date(2026, 9, 1)
    WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=period,
        ist_hours=Decimal('50.00'),
        soll_hours=Decimal('40.00'),
        paid_total_hours=Decimal('38.00'),
        saldo_cumulative=Decimal('12.00'),
        hourly_rate=Decimal('15.50'),
    )
    PayrollStatement.objects.create(
        worker=worker,
        period=period,
        gross_amount=Decimal('589.00'),
        net_amount=Decimal('589.00'),
        transferred_amount=Decimal('589.00'),
        source='lexware_pdf_bundle',
        raw_data=[{
            'kind': 'payslip',
            'compensation_type': 'hourly',
            'quantity': '38.00',
            'hourly_rate': '15.50',
            'gross_amount': '589.00',
            'net_amount': '589.00',
            'payout_amount': '589.00',
            'supplements': [],
        }],
    )

    response = auth_admin.get('/api/working-time/lexware-docx/?month=2026-09')
    assert response.status_code == 200
    assert '03_Lexware_Abgleich_2026-09.docx' in response['Content-Disposition']
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        document_xml = archive.read('word/document.xml').decode('utf-8')
    assert 'Lexware Abgleich' in document_xml
    assert 'Anna Becker' in document_xml


@pytest.mark.django_db
def test_lexware_pdf_bundle_imports_hourly_payslip_and_payment(
    auth_admin, worker_user
):
    import io

    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen import canvas

    def make_pdf(lines):
        buffer = io.BytesIO()
        doc = canvas.Canvas(buffer)
        y = 800
        for line in lines:
            doc.drawString(40, y, line)
            y -= 18
        doc.save()
        return buffer.getvalue()

    worker = worker_user.worker_profile
    period = date(2026, 9, 1)
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=period,
        ist_hours=Decimal('50.00'),
        soll_hours=Decimal('0.00'),
        paid_total_hours=Decimal('0.00'),
        saldo_cumulative=Decimal('50.00'),
        hourly_rate=Decimal('0.00'),
        gross_amount=Decimal('0.00'),
    )
    payslip = SimpleUploadedFile(
        'Lohnabrechnungen_2026-09.pdf',
        make_pdf([
            'Abrechnung für September 2026 - Anna Becker',
            'erstellt mit Lexware Seite 1 von 1',
            'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
            '14 01.01.1990 1 ohne',
            'Entgelt',
            'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
            'Lohn LSG 38,90 15,50 € 602,95 €',
            'Gesamtbrutto 602,95 €',
            'Netto 602,95 €',
            'Auszahlungsbetrag 602,95 €',
        ]),
        content_type='application/pdf',
    )
    payment = SimpleUploadedFile(
        'Zahlungsliste_2026-09.pdf',
        make_pdf([
            'Zahlungsliste September 2026',
            'Überweisung',
            'Mitarbeiter',
            'Empfänger Verwendungszweck IBAN Betrag',
            'Anna Becker Lohn & Gehalt September 2026 DE79 5085 2553 0117 5072 51 602,95',
        ]),
        content_type='application/pdf',
    )

    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-09', 'files': [payslip, payment]},
        format='multipart',
    )
    assert response.status_code == 200
    assert response.data['files'] == 2
    assert response.data['unmatched_count'] == 0

    statement = PayrollStatement.objects.get(worker=worker, period=period)
    assert statement.gross_amount == Decimal('602.95')
    assert statement.net_amount == Decimal('602.95')
    assert statement.transferred_amount == Decimal('602.95')
    assert statement.source == 'lexware_pdf_bundle'
    assert len(statement.raw_data) == 2

    record.refresh_from_db()
    assert record.paid_total_hours == Decimal('38.90')
    assert record.saldo_cumulative == Decimal('11.10')
    worker.refresh_from_db()
    assert worker.tariff_hourly_rate == Decimal('15.50')
    setting = worker.working_time_setting
    assert setting.hourly_rate == Decimal('15.50')


@pytest.mark.django_db
def test_lexware_salary_pdf_does_not_infer_paid_hours(auth_admin, worker_user):
    import io

    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer)
    for index, line in enumerate([
        'Abrechnung für September 2026 - Anna Becker',
        'erstellt mit Lexware Seite 1 von 1',
        'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
        '14 01.01.1990 1 ohne',
        'Entgelt',
        'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
        'Gehalt LSG 1,00 2.975,00 € 2.975,00 €',
        'Gesamtbrutto 2.975,00 €',
        'Netto 2.049,81 €',
        'Auszahlungsbetrag 2.099,81 €',
    ]):
        doc.drawString(40, 800 - index * 18, line)
    doc.save()

    worker = worker_user.worker_profile
    period = date(2026, 9, 1)
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=period,
        ist_hours=Decimal('96.98'),
        paid_total_hours=Decimal('0.00'),
        saldo_cumulative=Decimal('96.98'),
    )
    upload = SimpleUploadedFile(
        'Lohnabrechnungen_2026-09.pdf',
        buffer.getvalue(),
        content_type='application/pdf',
    )
    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-09', 'file': upload},
        format='multipart',
    )
    assert response.status_code == 200

    record.refresh_from_db()
    assert record.paid_total_hours == Decimal('0.00')
    statement = PayrollStatement.objects.get(worker=worker, period=period)
    assert statement.gross_amount == Decimal('2975.00')
    assert statement.net_amount == Decimal('2049.81')
    assert statement.transferred_amount is None
    assert statement.raw_data[0]['compensation_type'] == 'salary'
    assert statement.raw_data[0]['monthly_salary'] == '2975.00'


@pytest.mark.django_db
def test_each_employee_history_starts_with_first_authoritative_work_month(
    worker_user, second_worker, company, location, position
):
    tz = timezone.get_current_timezone()
    cases = [
        (worker_user.worker_profile, date(2026, 1, 5)),
        (second_worker, date(2026, 3, 5)),
    ]
    for worker, work_day in cases:
        start = timezone.make_aware(datetime.combine(work_day, time(8, 0)), tz)
        shift = Shift.objects.create(
            client=company,
            location=location,
            position=position,
            worker=worker,
            starts_at=start,
            ends_at=start + timedelta(hours=8),
            break_minutes=0,
            status=Shift.Status.CONFIRMED,
        )
        TimeEntry.objects.create(
            worker=worker,
            shift=shift,
            clock_in=start,
            clock_out=start + timedelta(hours=8),
            approved=True,
        )

    sync_working_time(date(2026, 1, 1), date(2026, 3, 31))

    first_worker_months = list(
        WorkingTimeAccountRecord.objects
        .filter(worker=worker_user.worker_profile)
        .order_by('year_month')
        .values_list('year_month', flat=True)
    )
    second_worker_months = list(
        WorkingTimeAccountRecord.objects
        .filter(worker=second_worker)
        .order_by('year_month')
        .values_list('year_month', flat=True)
    )

    assert first_worker_months[0] == date(2026, 1, 1)
    assert second_worker_months == [date(2026, 3, 1)]



@pytest.mark.django_db
def test_payroll_balance_skips_months_without_closed_attendance(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('100.00')
    worker.employment_type = 'vollzeit'
    worker.save(update_fields=['monthly_hours', 'employment_type', 'updated_at'])
    EmployeeMasterData.objects.create(
        worker=worker,
        data={'compensation_type': 'salary', 'monthly_salary': '2000.00'},
    )
    WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('100.00'),
        hourly_rate=Decimal('0.00'),
    )

    tz = timezone.get_current_timezone()
    for work_day in (date(2026, 1, 5), date(2026, 9, 5)):
        start = timezone.make_aware(datetime.combine(work_day, time(8, 0)), tz)
        shift = Shift.objects.create(
            client=company,
            location=location,
            position=position,
            worker=worker,
            starts_at=start,
            ends_at=start + timedelta(hours=8),
            break_minutes=0,
            status=Shift.Status.CONFIRMED,
        )
        TimeEntry.objects.create(
            worker=worker,
            shift=shift,
            clock_in=start,
            clock_out=start + timedelta(hours=8),
            approved=True,
        )

    # Simulate a stale row created by the old continuous-month rebuild logic.
    WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 2, 1),
        ist_hours=Decimal('0.00'),
        soll_hours=Decimal('100.00'),
        paid_total_hours=Decimal('0.00'),
        saldo_cumulative=Decimal('-192.00'),
        source='aplus_time_entries',
    )

    sync_working_time(date(2026, 1, 1), date(2026, 9, 30))

    months = list(
        WorkingTimeAccountRecord.objects
        .filter(worker=worker)
        .order_by('year_month')
        .values_list('year_month', flat=True)
    )
    assert months == [date(2026, 1, 1), date(2026, 9, 1)]

    january = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=date(2026, 1, 1))
    september = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=date(2026, 9, 1))
    assert january.saldo_cumulative == Decimal('-92.00')
    assert september.carryover_previous == Decimal('-92.00')
    assert september.saldo_cumulative == Decimal('-184.00')


@pytest.mark.django_db
def test_rebuild_uses_lexware_compensation_type_per_month(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('100.00')
    worker.employment_type = 'vollzeit'
    worker.save(update_fields=['monthly_hours', 'employment_type', 'updated_at'])
    EmployeeMasterData.objects.create(
        worker=worker,
        data={'compensation_type': 'salary', 'monthly_salary': '2000.00'},
    )
    WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('100.00'),
        hourly_rate=Decimal('20.00'),
    )

    tz = timezone.get_current_timezone()
    for work_day in (date(2026, 1, 5), date(2026, 9, 5)):
        start = timezone.make_aware(datetime.combine(work_day, time(8, 0)), tz)
        shift = Shift.objects.create(
            client=company,
            location=location,
            position=position,
            worker=worker,
            starts_at=start,
            ends_at=start + timedelta(hours=8),
            break_minutes=0,
            status=Shift.Status.CONFIRMED,
        )
        TimeEntry.objects.create(
            worker=worker,
            shift=shift,
            clock_in=start,
            clock_out=start + timedelta(hours=8),
            approved=True,
        )

    PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 1, 1),
        gross_amount=Decimal('160.00'),
        net_amount=Decimal('160.00'),
        transferred_amount=Decimal('160.00'),
        source='lexware_payslip_pdf',
        raw_data=[{
            'kind': 'payslip',
            'period': '2026-01',
            'compensation_type': 'hourly',
            'quantity': '8.00',
            'hourly_rate': '20.00',
            'gross_amount': '160.00',
            'net_amount': '160.00',
            'payout_amount': '160.00',
            'supplements': [],
        }],
    )
    PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 9, 1),
        gross_amount=Decimal('2000.00'),
        net_amount=Decimal('1500.00'),
        transferred_amount=Decimal('1500.00'),
        source='lexware_payslip_pdf',
        raw_data=[{
            'kind': 'payslip',
            'period': '2026-09',
            'compensation_type': 'salary',
            'quantity': '1.00',
            'monthly_salary': '2000.00',
            'gross_amount': '2000.00',
            'net_amount': '1500.00',
            'payout_amount': '1500.00',
            'supplements': [],
        }],
    )

    sync_working_time(
        date(2026, 1, 1),
        date(2026, 9, 30),
        include_inactive_workers=True,
        reset_carry=True,
    )

    january = WorkingTimeAccountRecord.objects.get(
        worker=worker,
        year_month=date(2026, 1, 1),
    )
    september = WorkingTimeAccountRecord.objects.get(
        worker=worker,
        year_month=date(2026, 9, 1),
    )
    assert january.paid_total_hours == Decimal('8.00')
    assert january.saldo_cumulative == Decimal('0.00')
    assert september.carryover_previous == Decimal('0.00')
    assert september.saldo_cumulative == Decimal('-92.00')


@pytest.mark.django_db
def test_scoped_payroll_rebuild_resets_prior_year_carry(
    auth_admin, worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('100.00')
    worker.employment_type = 'vollzeit'
    worker.save(update_fields=['monthly_hours', 'employment_type', 'updated_at'])
    EmployeeMasterData.objects.create(
        worker=worker,
        data={'compensation_type': 'salary', 'monthly_salary': '2000.00'},
    )
    WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('100.00'),
        hourly_rate=Decimal('0.00'),
    )
    WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2025, 12, 1),
        ist_hours=Decimal('0.00'),
        soll_hours=Decimal('100.00'),
        paid_total_hours=Decimal('0.00'),
        saldo_cumulative=Decimal('-500.00'),
        source='aplus_time_entries',
    )

    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime(2026, 9, 5, 8, 0), tz)
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        approved=True,
    )

    response = auth_admin.post(
        '/api/working-time/rebuild-all/',
        {'year': '2026'},
        format='json',
    )

    assert response.status_code == 200
    september = WorkingTimeAccountRecord.objects.get(
        worker=worker,
        year_month=date(2026, 9, 1),
    )
    assert september.carryover_previous == Decimal('0.00')
    assert september.saldo_cumulative == Decimal('-92.00')
    assert response.data['year'] == '2026'
    assert response.data['metadata']['reset_carry'] is True


@pytest.mark.django_db
def test_closed_month_keeps_historical_contract_and_rate_snapshot(
    worker_user, company, location, position
):
    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('38.00')
    worker.tariff_hourly_rate = Decimal('15.00')
    worker.employment_type = 'minijob'
    worker.save(update_fields=[
        'monthly_hours', 'tariff_hourly_rate', 'employment_type', 'updated_at'
    ])

    month = date(2026, 1, 1)
    start = timezone.make_aware(
        datetime.combine(month + timedelta(days=4), time(8, 0)),
        timezone.get_current_timezone(),
    )
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        approved=True,
    )

    sync_working_time(month, date(2026, 1, 31))
    record = WorkingTimeAccountRecord.objects.get(worker=worker, year_month=month)
    assert record.soll_hours == Decimal('38.00')
    assert record.hourly_rate == Decimal('15.00')
    assert record.employment_type_snapshot == 'minijob'

    worker.monthly_hours = Decimal('80.00')
    worker.tariff_hourly_rate = Decimal('20.00')
    worker.employment_type = 'teilzeit'
    worker.save(update_fields=[
        'monthly_hours', 'tariff_hourly_rate', 'employment_type', 'updated_at'
    ])
    setting = WorkingTimeSetting.objects.get(worker=worker)
    setting.monthly_limit = Decimal('80.00')
    setting.hourly_rate = Decimal('20.00')
    setting.save(update_fields=['monthly_limit', 'hourly_rate', 'updated_at'])

    sync_working_time(month, date(2026, 1, 31))
    record.refresh_from_db()
    assert record.soll_hours == Decimal('38.00')
    assert record.hourly_rate == Decimal('15.00')
    assert record.employment_type_snapshot == 'minijob'


@pytest.mark.django_db
def test_lexware_import_matches_inactive_former_employee(auth_admin, worker_user):
    from django.core.files.uploadedfile import SimpleUploadedFile

    worker = worker_user.worker_profile
    worker.active = False
    worker.save(update_fields=['active', 'updated_at'])
    period = date(2026, 4, 1)
    upload = SimpleUploadedFile(
        'bank.csv',
        b'Betrag;Datum;Verwendungszweck\n500,00;15.04.2026;Gehalt Anna Becker\n',
        content_type='text/csv',
    )
    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-04', 'file': upload},
        format='multipart',
    )
    assert response.status_code == 200
    assert response.data['employees'][0]['worker_id'] == str(worker.id)
    assert PayrollStatement.objects.get(worker=worker, period=period).transferred_amount == Decimal('500.00')


@pytest.mark.django_db
def test_lexware_import_rejects_ambiguous_employee_name(auth_admin, worker_user, second_worker):
    from django.core.files.uploadedfile import SimpleUploadedFile

    second_worker.user.first_name = 'Anna'
    second_worker.user.last_name = 'Becker'
    second_worker.user.save(update_fields=['first_name', 'last_name'])
    upload = SimpleUploadedFile(
        'bank.csv',
        b'Betrag;Datum;Verwendungszweck\n500,00;15.04.2026;Gehalt Anna Becker\n',
        content_type='text/csv',
    )
    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-04', 'file': upload},
        format='multipart',
    )
    assert response.status_code == 200
    assert response.data['employees'] == []
    assert response.data['unmatched_count'] == 1
    assert not PayrollStatement.objects.filter(period=date(2026, 4, 1)).exists()


@pytest.mark.django_db
def test_payroll_reconciliation_exposes_contract_warnings(worker_user):
    worker = worker_user.worker_profile
    worker.employment_type = 'minijob'
    worker.save(update_fields=['employment_type', 'updated_at'])
    EmployeeMasterData.objects.update_or_create(
        worker=worker,
        defaults={'data': {'compensation_type': 'hourly', 'hourly_rate': '16.00'}},
    )
    period = date(2026, 9, 1)
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=period,
        ist_hours=Decimal('38.00'),
        soll_hours=Decimal('38.00'),
        paid_total_hours=Decimal('38.00'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('16.00'),
        gross_amount=Decimal('2975.00'),
    )
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=period,
        gross_amount=Decimal('2975.00'),
        net_amount=Decimal('2049.81'),
        transferred_amount=Decimal('2099.81'),
        source='lexware_pdf_bundle',
        raw_data=[{
            'kind': 'payslip',
            'compensation_type': 'salary',
            'monthly_salary': '2975.00',
            'gross_amount': '2975.00',
            'net_amount': '2049.81',
            'payout_amount': '2099.81',
            'supplements': [],
        }],
    )

    from core.working_time import record_dict
    data = record_dict(record, statement)

    assert data['minijob_warning'] is True
    assert any('Vergütungsart stimmt nicht überein' in item for item in data['contract_issues'])
    assert any('Minijob prüfen' in item for item in data['contract_issues'])
    assert data['reconciliation_status'] == 'ABWEICHUNG'


@pytest.mark.django_db
def test_payroll_reconciliation_compares_night_weekend_hours(worker_user):
    worker = worker_user.worker_profile
    worker.employment_type = 'teilzeit'
    worker.save(update_fields=['employment_type', 'updated_at'])
    EmployeeMasterData.objects.update_or_create(
        worker=worker,
        defaults={'data': {'compensation_type': 'hourly', 'hourly_rate': '16.00'}},
    )
    period = date(2026, 9, 1)
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=period,
        ist_hours=Decimal('8.00'),
        soll_hours=Decimal('8.00'),
        paid_total_hours=Decimal('8.00'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('16.00'),
        gross_amount=Decimal('128.00'),
        raw_entries=[{
            'worked_minutes': 480,
            'break_minutes': 0,
            'night_minutes': 120,
            'saturday_minutes': 60,
            'sunday_minutes': 0,
            'night_surcharge_amount': '8.00',
            'saturday_surcharge_amount': '4.00',
            'sunday_surcharge_amount': '0.00',
        }],
    )
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=period,
        gross_amount=Decimal('140.00'),
        net_amount=Decimal('140.00'),
        transferred_amount=Decimal('140.00'),
        source='lexware_pdf_bundle',
        raw_data=[{
            'kind': 'payslip',
            'compensation_type': 'hourly',
            'quantity': '8.00',
            'hourly_rate': '16.00',
            'payout_amount': '140.00',
            'supplements': [
                {'label': 'Nachtzuschlag', 'hours': '2.00', 'percent': '25', 'amount': '8.00'},
                {'label': 'Samstagszuschlag', 'hours': '1.00', 'percent': '25', 'amount': '4.00'},
            ],
        }],
    )

    from core.working_time import record_dict
    data = record_dict(record, statement)

    assert data['surcharge_reconciliation']['overall'] == 'MATCH'
    assert data['surcharge_reconciliation']['night']['status'] == 'MATCH'
    assert data['surcharge_reconciliation']['saturday']['status'] == 'MATCH'
    assert data['surcharge_reconciliation']['sunday']['status'] == 'MATCH'
    assert data['reconciliation_status'] == 'MATCH'

    statement.raw_data = [{
        'kind': 'payslip',
        'compensation_type': 'hourly',
        'quantity': '8.00',
        'hourly_rate': '16.00',
        'payout_amount': '140.00',
        'supplements': [
            {'label': 'Nachtzuschlag', 'hours': '1.00', 'percent': '25', 'amount': '4.00'},
            {'label': 'Samstagszuschlag', 'hours': '1.00', 'percent': '25', 'amount': '4.00'},
        ],
    }]
    statement.save(update_fields=['raw_data', 'updated_at'])

    changed = record_dict(record, statement)
    assert changed['surcharge_reconciliation']['night']['status'] == 'ABWEICHUNG'
    assert changed['reconciliation_status'] == 'ABWEICHUNG'


@pytest.mark.django_db
def test_worktime_settings_persists_surcharge_percentages(auth_admin, worker_user):
    worker = worker_user.worker_profile
    response = auth_admin.post(
        '/api/working-time/settings/',
        {
            'employees': [{
                'worker_id': str(worker.id),
                'monthly_limit': '80.00',
                'hourly_rate': '16.00',
                'night_surcharge_percent': '25.00',
                'saturday_surcharge_percent': '20.00',
                'sunday_surcharge_percent': '50.00',
                'active': True,
                'excluded': False,
            }],
        },
        format='json',
    )

    assert response.status_code == 200
    setting = WorkingTimeSetting.objects.get(worker=worker)
    assert setting.monthly_limit == Decimal('80.00')
    assert setting.hourly_rate == Decimal('16.00')
    assert setting.night_surcharge_percent == Decimal('25.00')
    assert setting.saturday_surcharge_percent == Decimal('20.00')
    assert setting.sunday_surcharge_percent == Decimal('50.00')


@pytest.mark.django_db
def test_lexware_sepa_xml_does_not_double_count_payment_list(auth_admin, worker_user):
    import io

    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen import canvas
    from core.models import Document

    def make_pdf(lines):
        buffer = io.BytesIO()
        doc = canvas.Canvas(buffer)
        y = 800
        for line in lines:
            doc.drawString(40, y, line)
            y -= 18
        doc.save()
        return buffer.getvalue()

    worker = worker_user.worker_profile
    period = date(2026, 9, 1)
    payment = SimpleUploadedFile(
        'Zahlungsliste_2026-09.pdf',
        make_pdf([
            'Zahlungsliste September 2026',
            'Überweisung',
            'Mitarbeiter',
            'Empfänger Verwendungszweck IBAN Betrag',
            'Anna Becker Lohn & Gehalt September 2026 DE79 5085 2553 0117 5072 51 602,95',
        ]),
        content_type='application/pdf',
    )
    sepa = SimpleUploadedFile(
        'SEPA_Ueberweisungstraeger_2026-09.xml',
        b'''<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:pain.001.001.09">
  <CstmrCdtTrfInitn>
    <PmtInf>
      <ReqdExctnDt><Dt>2026-10-08</Dt></ReqdExctnDt>
      <CdtTrfTxInf>
        <Amt><InstdAmt Ccy="EUR">602.95</InstdAmt></Amt>
        <Cdtr><Nm>Anna Becker</Nm></Cdtr>
        <CdtrAcct><Id><IBAN>DE79508525530117507251</IBAN></Id></CdtrAcct>
        <RmtInf><Ustrd>Lohn und Gehalt September 2026</Ustrd></RmtInf>
      </CdtTrfTxInf>
    </PmtInf>
  </CstmrCdtTrfInitn>
</Document>''',
        content_type='application/xml',
    )

    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-09', 'files': [payment, sepa]},
        format='multipart',
    )

    assert response.status_code == 200
    statement = PayrollStatement.objects.get(worker=worker, period=period)
    assert statement.transferred_amount == Decimal('602.95')
    assert statement.payment_date is None
    assert statement.source == 'lexware_sepa_xml'
    assert {row.get('source_type') for row in statement.raw_data} == {
        'lexware_zahlungsliste_pdf',
        'lexware_sepa_xml',
    }
    assert len(response.data['archived_documents']) == 2
    assert Document.objects.filter(folder='payroll', visibility='admin').count() >= 2


@pytest.mark.django_db
def test_cross_month_correction_payslip_updates_original_period(auth_admin, worker_user):
    import io

    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer)

    def page(lines):
        y = 800
        for line in lines:
            doc.drawString(40, y, line)
            y -= 18
        doc.showPage()

    page([
        'Korrekturabrechnung für Juni 2026 - Anna Becker',
        'erstellt mit Lexware Seite 1 von 1',
        'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
        '14 01.01.1990 1 ohne',
        'Entgelt',
        'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
        'Lohn LSG 10,00 15,50 € 155,00 €',
        'Gesamtbrutto 155,00 €',
        'Netto 140,00 €',
        'Persönliche Be-/Abzüge',
        'Vorschuss aus Überzahlung 15,00 €',
        'Bereits abgerechnete Auszahlung -155,00 €',
        'Auszahlungsbetrag 0,00 €',
    ])
    page([
        'Abrechnung für Juli 2026 - Anna Becker',
        'erstellt mit Lexware Seite 1 von 1',
        'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
        '14 01.01.1990 1 ohne',
        'Entgelt',
        'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
        'Lohn LSG 20,00 15,50 € 310,00 €',
        'Gesamtbrutto 310,00 €',
        'Netto 310,00 €',
        'Auszahlungsbetrag 310,00 €',
    ])
    doc.save()

    upload = SimpleUploadedFile(
        'Lohnabrechnungen_2026-07.pdf',
        buffer.getvalue(),
        content_type='application/pdf',
    )

    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-07', 'file': upload},
        format='multipart',
    )

    assert response.status_code == 200
    june = PayrollStatement.objects.get(
        worker=worker_user.worker_profile,
        period=date(2026, 6, 1),
    )
    july = PayrollStatement.objects.get(
        worker=worker_user.worker_profile,
        period=date(2026, 7, 1),
    )
    assert june.gross_amount == Decimal('155.00')
    assert june.net_amount == Decimal('140.00')
    assert june.raw_data[0]['is_correction'] is True
    assert june.raw_data[0]['period'] == '2026-06'
    assert june.raw_data[0]['personal_adjustments'] == [
        {'label': 'Vorschuss aus Überzahlung', 'amount': '15.00'},
        {'label': 'Bereits abgerechnete Auszahlung', 'amount': '-155.00'},
    ]
    assert july.gross_amount == Decimal('310.00')
    assert july.net_amount == Decimal('310.00')
    assert july.raw_data[0]['is_correction'] is False
    assert july.raw_data[0]['period'] == '2026-07'


def test_lexware_payment_list_accepts_und_and_ampersand_wording():
    import io

    from reportlab.pdfgen import canvas
    from core.lexware_pdf import parse_payment_list

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer)
    doc.drawString(40, 800, 'Anna Becker Lohn und Gehalt März 2026 DE79 5085 2553 0117 5072 51 123,45')
    doc.drawString(40, 780, 'Max Muster Lohn & Gehalt April 2026 DE36 5055 0020 0105 2038 61 234,56')
    doc.save()

    rows = parse_payment_list(buffer.getvalue())

    assert [(row['employee_name'], row['amount']) for row in rows] == [
        ('Anna Becker', '123.45'),
        ('Max Muster', '234.56'),
    ]


@pytest.mark.django_db
def test_full_year_zip_import_splits_months_and_archives_unknown_pdfs(auth_admin, worker_user):
    import io
    import zipfile

    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen import canvas
    from core.models import Document

    def make_pdf(lines):
        buffer = io.BytesIO()
        doc = canvas.Canvas(buffer)
        y = 800
        for line in lines:
            doc.drawString(40, y, line)
            y -= 18
        doc.save()
        return buffer.getvalue()

    january = make_pdf([
        'Zahlungsliste Januar 2026',
        'Anna Becker Lohn und Gehalt Januar 2026 DE79 5085 2553 0117 5072 51 123,45',
    ])
    february = make_pdf([
        'Zahlungsliste Februar 2026',
        'Anna Becker Lohn & Gehalt Februar 2026 DE79 5085 2553 0117 5072 51 234,56',
    ])
    journal = make_pdf([
        'Lohnjournal Januar 2026',
        'A+ Solution GmbH',
        'Gesamtsumme 123,45',
    ])
    prior_year_correction = make_pdf([
        'Korrekturabrechnung für Dezember 2025 - Anna Becker',
        'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
        'MA-001 01.01.1990 1 ohne',
        'Entgelt',
        'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
        'Lohn LSG 10,00 15,50 € 155,00 €',
        'Gesamtbrutto 155,00 €',
        'Netto 155,00 €',
        'Bereits abgerechnete Auszahlung -155,00 €',
        'Auszahlungsbetrag 0,00 €',
    ])

    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('01/Zahlungsliste_2026-01.pdf', january)
        archive.writestr('02/Zahlungsliste_2026-02.pdf', february)
        archive.writestr('01/2026-01_Lohnjournal.pdf', journal)
        archive.writestr('02/Lohnabrechnungen_2026-02.pdf', prior_year_correction)

    upload = SimpleUploadedFile(
        'Aplus_Lexware_2026.zip',
        bundle.getvalue(),
        content_type='application/zip',
    )

    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': 'auto', 'year': '2026', 'file': upload},
        format='multipart',
    )

    assert response.status_code == 200
    worker = worker_user.worker_profile
    january_statement = PayrollStatement.objects.get(worker=worker, period=date(2026, 1, 1))
    february_statement = PayrollStatement.objects.get(worker=worker, period=date(2026, 2, 1))
    assert january_statement.transferred_amount == Decimal('123.45')
    assert february_statement.transferred_amount == Decimal('234.56')
    assert response.data['detected_periods'] == ['2026-01', '2026-02']
    assert '2026-01_Lohnjournal.pdf' in response.data['archived_only']
    assert not PayrollStatement.objects.filter(worker=worker, period=date(2025, 12, 1)).exists()
    assert response.data['skipped_historical_corrections'] == [{
        'file': 'Lohnabrechnungen_2026-02.pdf',
        'employee_name': 'Anna Becker',
        'period': '2025-12',
    }]
    assert Document.objects.filter(folder='payroll', visibility='admin').count() == 4


@pytest.mark.django_db
def test_zero_expected_surcharge_does_not_create_false_lexware_warning(worker_user):
    from core.working_time import record_dict

    worker = worker_user.worker_profile
    EmployeeMasterData.objects.create(
        worker=worker,
        data={
            'compensation_type': 'salary',
            'monthly_salary': '2975.00',
            'lexware_latest_payroll_period': '2026-09',
        },
    )
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('96.98'),
        soll_hours=Decimal('166.83'),
        paid_total_hours=Decimal('0.00'),
        saldo_cumulative=Decimal('-69.85'),
        hourly_rate=Decimal('0.00'),
        gross_amount=Decimal('2975.00'),
        raw_entries=[{
            'night_minutes': 14,
            'saturday_minutes': 465,
            'sunday_minutes': 0,
            'night_surcharge_amount': '0.00',
            'saturday_surcharge_amount': '0.00',
            'sunday_surcharge_amount': '0.00',
        }],
    )
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 9, 1),
        gross_amount=Decimal('2975.00'),
        net_amount=Decimal('2049.81'),
        transferred_amount=Decimal('2099.81'),
        source='lexware_payslip_pdf',
        raw_data=[{
            'kind': 'payslip',
            'period': '2026-09',
            'compensation_type': 'salary',
            'monthly_salary': '2975.00',
            'gross_amount': '2975.00',
            'net_amount': '2049.81',
            'payout_amount': '2099.81',
            'supplements': [],
        }],
    )

    data = record_dict(record, statement)

    assert data['surcharge_reconciliation']['night']['status'] == 'MATCH'
    assert data['surcharge_reconciliation']['saturday']['status'] == 'MATCH'
    assert not any('Zuschlag nicht nachgewiesen' in issue for issue in data['reconciliation_issues'])


@pytest.mark.django_db
def test_historical_lexware_type_does_not_compare_to_current_master(worker_user):
    from core.working_time import record_dict

    worker = worker_user.worker_profile
    EmployeeMasterData.objects.create(
        worker=worker,
        data={
            'compensation_type': 'salary',
            'monthly_salary': '2975.00',
            'lexware_latest_payroll_period': '2026-09',
        },
    )
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 6, 1),
        ist_hours=Decimal('38.90'),
        soll_hours=Decimal('38.90'),
        paid_total_hours=Decimal('38.90'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('15.50'),
        gross_amount=Decimal('602.95'),
    )
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 6, 1),
        gross_amount=Decimal('614.59'),
        net_amount=Decimal('614.59'),
        transferred_amount=Decimal('614.59'),
        source='lexware_pdf_bundle',
        raw_data=[{
            'kind': 'payslip',
            'period': '2026-06',
            'compensation_type': 'hourly',
            'quantity': '38.90',
            'hourly_rate': '15.50',
            'gross_amount': '614.59',
            'net_amount': '614.59',
            'payout_amount': '614.59',
            'supplements': [{
                'label': 'Nachtzuschlag 25% (steuerfrei)',
                'hours': '3.00',
                'hourly_rate': '15.50',
                'percent': '25.00',
                'amount': '11.64',
            }],
        }],
    )

    data = record_dict(record, statement)

    assert not any('Vergütungsart stimmt nicht überein' in issue for issue in data['contract_issues'])


@pytest.mark.django_db
def test_older_lexware_import_does_not_regress_current_master_rate(auth_admin, worker_user):
    import io

    from django.core.files.uploadedfile import SimpleUploadedFile
    from reportlab.pdfgen import canvas

    def make_pdf(lines):
        buffer = io.BytesIO()
        doc = canvas.Canvas(buffer)
        y = 800
        for line in lines:
            doc.drawString(40, y, line)
            y -= 18
        doc.save()
        return buffer.getvalue()

    worker = worker_user.worker_profile
    worker.tariff_hourly_rate = Decimal('20.00')
    worker.save(update_fields=['tariff_hourly_rate', 'updated_at'])

    september = SimpleUploadedFile(
        'Lohnabrechnungen_2026-09.pdf',
        make_pdf([
            'Abrechnung für September 2026 - Anna Becker',
            'erstellt mit Lexware Seite 1 von 1',
            'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
            '14 01.01.1990 1 ohne',
            'Entgelt',
            'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
            'Gehalt LSG 1,00 2.975,00 € 2.975,00 €',
            'Gesamtbrutto 2.975,00 €',
            'Netto 2.049,81 €',
            'Auszahlungsbetrag 2.099,81 €',
        ]),
        content_type='application/pdf',
    )
    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-09', 'file': september},
        format='multipart',
    )
    assert response.status_code == 200

    january = SimpleUploadedFile(
        'Lohnabrechnungen_2026-01.pdf',
        make_pdf([
            'Abrechnung für Januar 2026 - Anna Becker',
            'erstellt mit Lexware Seite 1 von 1',
            'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
            '14 01.01.1990 1 ohne',
            'Entgelt',
            'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
            'Lohn LSG 18,07 15,50 € 280,09 €',
            'Gesamtbrutto 280,09 €',
            'Netto 280,09 €',
            'Auszahlungsbetrag 280,09 €',
        ]),
        content_type='application/pdf',
    )
    response = auth_admin.post(
        '/api/working-time/lexware-import/',
        {'period': '2026-01', 'file': january},
        format='multipart',
    )
    assert response.status_code == 200

    master = EmployeeMasterData.objects.get(worker=worker)
    worker.refresh_from_db()
    assert master.data['lexware_latest_payroll_period'] == '2026-09'
    assert master.data['compensation_type'] == 'salary'
    assert master.data['lexware_monthly_salary'] == '2975.00'
    assert worker.tariff_hourly_rate == Decimal('20.00')


@pytest.mark.django_db
def test_zero_payout_correction_uses_original_monthly_payout_for_reconciliation(worker_user):
    from core.working_time import record_dict

    worker = worker_user.worker_profile
    EmployeeMasterData.objects.create(
        worker=worker,
        data={
            'compensation_type': 'salary',
            'monthly_salary': '2975.00',
            'lexware_latest_payroll_period': '2026-09',
        },
    )
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 7, 1),
        ist_hours=Decimal('167.50'),
        soll_hours=Decimal('167.50'),
        paid_total_hours=Decimal('167.50'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('17.12'),
        gross_amount=Decimal('2867.60'),
        raw_entries=[{
            'night_minutes': 300,
            'saturday_minutes': 0,
            'sunday_minutes': 0,
            'night_surcharge_amount': '21.40',
            'saturday_surcharge_amount': '0.00',
            'sunday_surcharge_amount': '0.00',
        }],
    )
    common = {
        'kind': 'payslip',
        'period': '2026-07',
        'compensation_type': 'hourly',
        'person_group': '101',
        'quantity': '167.50',
        'hourly_rate': '17.12',
        'gross_amount': '2889.00',
        'supplements': [{
            'label': 'Nachtzuschlag 25% (steuerfrei)',
            'hours': '5.00',
            'hourly_rate': '17.12',
            'percent': '25.00',
            'amount': '21.40',
        }],
    }
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 7, 1),
        gross_amount=Decimal('2889.00'),
        net_amount=Decimal('2006.40'),
        transferred_amount=Decimal('2029.71'),
        source='lexware_pdf_bundle',
        raw_data=[
            {
                **common,
                'is_correction': False,
                'net_amount': '2029.71',
                'payout_amount': '2029.71',
            },
            {
                **common,
                'is_correction': True,
                'net_amount': '2006.40',
                'payout_amount': '0.00',
                'personal_adjustments': [
                    {'label': 'Vorschuss aus Überzahlung', 'amount': '23.31'},
                    {'label': 'Bereits abgerechnete Auszahlung', 'amount': '-2029.71'},
                ],
            },
        ],
    )

    data = record_dict(record, statement)

    assert data['payroll_statement']['net_amount'] == '2006.40'
    assert data['payroll_statement']['lexware_payout_amount'] == '2029.71'
    assert data['payroll_statement']['lexware_correction_payout_amount'] == '0.00'
    assert data['reconciliation_status'] == 'MATCH'


def test_lexware_payslip_parser_reads_person_group():
    import io

    from reportlab.pdfgen import canvas
    from core.lexware_pdf import parse_payslips

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer)
    for index, line in enumerate([
        'Abrechnung für Januar 2026 - Anna Becker',
        'Personal-Nr. Geburtsdatum Steuerklasse Konfession',
        '14 01.01.1990 - ohne',
        'Pers.-Grp. Beitragsgruppe Eintritt Austritt',
        '109 6500 01.10.2024 -',
        'Entgelt',
        'Bezeichnung Kennz Menge Faktor Prozentsatz Betrag',
        'Lohn LSG 18,07 15,50 € 280,09 €',
        'Gesamtbrutto 280,09 €',
        'Netto 280,09 €',
        'Auszahlungsbetrag 280,09 €',
    ]):
        doc.drawString(40, 800 - index * 18, line)
    doc.save()

    rows = parse_payslips(buffer.getvalue())

    assert len(rows) == 1
    assert rows[0]['person_group'] == '109'


@pytest.mark.django_db
def test_historical_person_group_109_displays_as_minijob(worker_user):
    from core.working_time import record_dict

    worker = worker_user.worker_profile
    worker.employment_type = 'vollzeit'
    worker.save(update_fields=['employment_type', 'updated_at'])
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 1, 1),
        ist_hours=Decimal('18.07'),
        soll_hours=Decimal('38.00'),
        paid_total_hours=Decimal('18.07'),
        employment_type_snapshot='vollzeit',
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('15.50'),
        gross_amount=Decimal('280.09'),
    )
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 1, 1),
        gross_amount=Decimal('280.09'),
        net_amount=Decimal('280.09'),
        transferred_amount=Decimal('280.09'),
        source='lexware_pdf_bundle',
        raw_data=[{
            'kind': 'payslip',
            'period': '2026-01',
            'person_group': '109',
            'compensation_type': 'hourly',
            'quantity': '18.07',
            'hourly_rate': '15.50',
            'gross_amount': '280.09',
            'net_amount': '280.09',
            'payout_amount': '280.09',
            'supplements': [],
        }],
    )

    data = record_dict(record, statement)

    assert data['employment_type'] == 'minijob'


@pytest.mark.django_db
def test_open_current_payroll_month_does_not_accrue_full_salary_deficit(
    worker_user, company, location, position
):
    from core.working_time import record_dict

    worker = worker_user.worker_profile
    worker.monthly_hours = Decimal('100.00')
    worker.employment_type = 'vollzeit'
    worker.save(update_fields=['monthly_hours', 'employment_type', 'updated_at'])
    EmployeeMasterData.objects.create(
        worker=worker,
        data={'compensation_type': 'salary', 'monthly_salary': '2000.00'},
    )
    WorkingTimeSetting.objects.create(
        worker=worker,
        monthly_limit=Decimal('100.00'),
        hourly_rate=Decimal('0.00'),
    )

    today = timezone.localdate()
    work_day = today.replace(day=max(1, min(today.day, 5)))
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(work_day, time(8, 0)), tz)
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=0,
        status=Shift.Status.CONFIRMED,
    )
    TimeEntry.objects.create(
        worker=worker,
        shift=shift,
        clock_in=start,
        clock_out=start + timedelta(hours=8),
        approved=False,
    )

    month_start = today.replace(day=1)
    sync_working_time(month_start, today, reset_carry=True)

    record = WorkingTimeAccountRecord.objects.get(
        worker=worker,
        year_month=month_start,
    )
    data = record_dict(record)

    assert record.saldo_cumulative == Decimal('0.00')
    assert record.gross_amount == Decimal('0.00')
    assert data['is_open_month'] is True
    assert data['balance_basis'] == 'open_month'
    assert data['monthly_balance_hours'] == '0.00'
    assert data['saldo_cumulative'] == '0.00'
    assert data['reconciliation_status'] == 'LAUFEND'


def test_lexware_u1_parser_extracts_sickness_evidence():
    import io

    from reportlab.pdfgen import canvas
    from core.lexware_pdf import parse_lexware_pdf

    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer)
    lines = [
        'Erstattungsantrag U1-Krankheit September 2026',
        'A+ Solution GmbH, Carl-Sonnenschein Strasse 57, 65936 Frankfurt am Main',
        'Arina Martynko, PNr: 6, Zwingenberger Strasse 3, 68519 Viernheim',
        'Erstattungszeitraum von Erstattungszeitraum bis Erstattungsbetrag',
        '15.09.2026 16.09.2026 73,05 €',
        'Entgelt 16,19 € Stundenlohn',
        'Ausfallzeit 5,64 Stunden',
        'Arbeitszeit 19,77 h wöchentlich / 2,82 h täglich',
        'Fortgezahltes Entgelt 91,31 €',
        'Erstattungssatz 80 %',
    ]
    for index, line in enumerate(lines):
        doc.drawString(40, 800 - index * 18, line)
    doc.save()

    document_type, rows = parse_lexware_pdf(buffer.getvalue())

    assert document_type == 'u1'
    assert len(rows) == 1
    assert rows[0]['employee_name'] == 'Arina Martynko'
    assert rows[0]['period'] == '2026-09'
    assert rows[0]['date_from'] == '15.09.2026'
    assert rows[0]['date_to'] == '16.09.2026'
    assert rows[0]['absence_hours'] == '5.64'
    assert rows[0]['continued_pay'] == '91.31'
    assert rows[0]['reimbursement_amount'] == '73.05'


@pytest.mark.django_db
def test_absence_summary_is_visible_without_changing_saldo(worker_user):
    from core.working_time import absence_summary_map, record_dict

    worker = worker_user.worker_profile
    TimeOffRequest.objects.create(
        worker=worker,
        starts_on=date(2026, 9, 15),
        ends_on=date(2026, 9, 16),
        reason='Krankheit',
        status=TimeOffRequest.Status.APPROVED,
    )
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('20.00'),
        soll_hours=Decimal('20.00'),
        paid_total_hours=Decimal('20.00'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('15.50'),
        gross_amount=Decimal('310.00'),
    )
    summary = absence_summary_map([record])
    data = record_dict(
        record,
        absence_summary=summary[(str(worker.id), date(2026, 9, 1))],
    )

    assert data['absence_days'] == 2
    assert data['sick_days'] == 2
    assert data['vacation_days'] == 0
    assert data['saldo_cumulative'] == '0.00'


@pytest.mark.django_db
def test_lexware_readiness_marks_complete_closed_month_package():
    from django.core.files.uploadedfile import SimpleUploadedFile
    from core.payroll_views import lexware_readiness_payload

    for filename in (
        'Lohnabrechnungen_2026-01.pdf',
        'Zahlungsliste_2026-01.pdf',
        '2026-01_Lohnjournal.pdf',
        'SEPA_Ueberweisungstraeger_2026-01.xml',
        'Meldebescheinigungen_2026-01.pdf',
        'Lohnkonto.pdf',
        'Lohnkonto-UV_2026.pdf',
    ):
        Document.objects.create(
            title=f'Lexware 2026-01 · {filename}' if '2026-01' in filename else f'Lexware 2026 · {filename}',
            file=SimpleUploadedFile(filename, b'test'),
            folder=Document.Folder.PAYROLL,
            visibility=Document.Visibility.ADMIN,
        )

    readiness = lexware_readiness_payload(2026)

    january = next(item for item in readiness['months'] if item['period'] == '2026-01')
    assert january['core_complete'] is True
    assert readiness['annual_documents']['lohnkonto'] is True
    assert readiness['annual_documents']['lohnkonto_uv'] is True


@pytest.mark.django_db
def test_payroll_audit_docx_generates_with_absence_columns(worker_user):
    from core.working_time import payroll_audit_docx

    worker = worker_user.worker_profile
    WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('20.00'),
        soll_hours=Decimal('20.00'),
        paid_total_hours=Decimal('20.00'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('15.50'),
        gross_amount=Decimal('310.00'),
    )
    payload = payroll_audit_docx(
        WorkingTimeAccountRecord.objects.filter(worker=worker, year_month__year=2026),
        2026,
        {
            'core_documents_present': 45,
            'core_documents_expected': 45,
            'complete_months': 9,
            'expected_months': 9,
            'annual_complete': True,
            'no_additional_import_required': True,
            'months': [],
        },
    )

    assert payload[:2] == b'PK'


@pytest.mark.django_db
def test_payroll_excel_contains_filterable_daily_evidence_and_notes(
    worker_user, company, location, position
):
    import io

    from openpyxl import load_workbook
    from core.working_time import export_xlsx

    worker = worker_user.worker_profile
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime(2026, 9, 5, 8, 0), tz)
    shift = Shift.objects.create(
        client=company,
        location=location,
        position=position,
        worker=worker,
        starts_at=start,
        ends_at=start + timedelta(hours=8),
        break_minutes=30,
        status=Shift.Status.CONFIRMED,
        notes='Event 4711 Eingang West',
    )
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('7.50'),
        soll_hours=Decimal('7.50'),
        paid_total_hours=Decimal('7.50'),
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('15.50'),
        gross_amount=Decimal('116.25'),
        raw_entries=[{
            'id': 'entry-1',
            'shift_id': str(shift.id),
            'client_name': company.name,
            'location_name': location.name,
            'position_name': position.name,
            'planned_start': start.isoformat(),
            'planned_end': (start + timedelta(hours=8)).isoformat(),
            'local_clock_in': start.isoformat(),
            'local_clock_out': (start + timedelta(hours=8)).isoformat(),
            'break_minutes': 30,
            'worked_minutes': 450,
            'night_minutes': 0,
            'saturday_minutes': 0,
            'sunday_minutes': 0,
        }],
    )

    response = export_xlsx(
        WorkingTimeAccountRecord.objects.filter(pk=record.pk)
    )
    workbook = load_workbook(io.BytesIO(response.content))
    assert 'Tagesnachweise' in workbook.sheetnames
    sheet = workbook['Tagesnachweise']
    headers = [cell.value for cell in sheet[1]]
    assert headers[:6] == ['Mitarbeiter', 'Monat', 'Datum', 'Kunde', 'Ort', 'Service']
    assert 'Notiz' in headers
    note_column = headers.index('Notiz') + 1
    assert sheet.cell(row=2, column=note_column).value == 'Event 4711 Eingang West'
    assert sheet.auto_filter.ref


@pytest.mark.django_db
def test_person_group_997_displays_as_managing_director(worker_user):
    from core.working_time import record_dict

    worker = worker_user.worker_profile
    worker.employment_type = 'minijob'
    worker.save(update_fields=['employment_type', 'updated_at'])
    record = WorkingTimeAccountRecord.objects.create(
        worker=worker,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('0.00'),
        soll_hours=Decimal('0.00'),
        paid_total_hours=Decimal('0.00'),
        employment_type_snapshot='minijob',
        saldo_cumulative=Decimal('0.00'),
        hourly_rate=Decimal('0.00'),
        gross_amount=Decimal('0.00'),
    )
    statement = PayrollStatement.objects.create(
        worker=worker,
        period=date(2026, 9, 1),
        gross_amount=Decimal('2800.00'),
        net_amount=Decimal('2523.67'),
        transferred_amount=Decimal('2523.67'),
        source='lexware_pdf_bundle',
        raw_data=[{
            'kind': 'payslip',
            'period': '2026-09',
            'person_group': '997',
            'compensation_type': 'salary',
            'monthly_salary': '2500.00',
            'gross_amount': '2800.00',
            'net_amount': '2523.67',
            'payout_amount': '2523.67',
            'supplements': [],
        }],
    )

    data = record_dict(record, statement)

    assert data['employment_type'] == 'geschaeftsfuehrer'
    assert data['minijob_limit'] is None
    assert data['payroll_statement']['gross_amount'] == '2800.00'
