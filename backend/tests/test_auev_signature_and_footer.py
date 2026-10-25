import io
import zipfile
from datetime import date

from django.test import SimpleTestCase
from docx import Document
from PIL import Image, ImageDraw

from core.auev_builder import _classic_docx, _new_docx, _normalise_provider_signature_image, _week_footer


def _signature_png():
    buffer = io.BytesIO()
    Image.new('RGBA', (90, 42), (0, 0, 0, 255)).save(buffer, 'PNG')
    return buffer.getvalue()


def _data(signature):
    return {
        'sequence_number': 415,
        'week_footer': '42-43',
        'provider_signature': signature,
        'signature_date_de': '15.10.2026',
        'first_shift_date_de': '17.10.2026',
        'last_shift_date_de': '24.10.2026',
        'settings': {
            'client_contract_text': 'Marthas Finest GmbH, Frankfurt',
            'permit_date_de': '20.04.2025',
            'framework_date_de': '26.08.2024',
            'effective_date_de': '17.10.2026',
            'intended_activity': 'Servicekraft',
        },
        'rows': [{
            'name_birth': 'Mustermann, Max, 01.01.2000',
            'start': '08:30', 'end': '15:30', 'date': '17.10.2026',
            'activity': 'Servicekraft',
        }],
    }


class AuevSignatureAndFooterTests(SimpleTestCase):
    def test_footer_uses_iso_week_values_and_not_document_sequence(self):
        self.assertEqual(_week_footer([date(2026, 10, 17), date(2026, 10, 24)]), '42-43')
        self.assertEqual(_week_footer([date(2026, 12, 31), date(2027, 1, 4)]), '53/2026-1/2027')

    def test_both_templates_embed_signature_and_display_week_footer(self):
        for build in (_classic_docx, _new_docx):
            with self.subTest(template=build.__name__):
                docx_bytes = build(_data(_signature_png()))
                rendered = Document(io.BytesIO(docx_bytes))
                self.assertEqual(rendered.sections[0].footer.paragraphs[0].text, '42-43')
                self.assertNotIn('415', rendered.sections[0].footer.paragraphs[0].text)
                with zipfile.ZipFile(io.BytesIO(docx_bytes)) as contents:
                    self.assertTrue(any(path.startswith('word/media/') for path in contents.namelist()))

    def test_unapproved_signature_is_not_embedded(self):
        for build in (_classic_docx, _new_docx):
            with self.subTest(template=build.__name__):
                docx_bytes = build(_data(None))
                with zipfile.ZipFile(io.BytesIO(docx_bytes)) as contents:
                    self.assertFalse(any(path.startswith('word/media/') for path in contents.namelist()))


    def test_faint_signature_is_darken_and_trimmed_without_changing_approved_source(self):
        canvas = Image.new('RGBA', (620, 220), (255, 255, 255, 255))
        ink = ImageDraw.Draw(canvas)
        ink.line((170, 110, 450, 85, 500, 105), fill=(222, 222, 222, 255), width=4)
        source = io.BytesIO()
        canvas.save(source, format='PNG')
        approved_bytes = source.getvalue()

        display_bytes, size = _normalise_provider_signature_image(approved_bytes)
        self.assertEqual(source.getvalue(), approved_bytes)
        self.assertLess(size[0], 400)
        self.assertLess(size[1], 80)
        with Image.open(io.BytesIO(display_bytes)) as rendered:
            self.assertEqual(rendered.mode, 'RGBA')
            alpha = rendered.getchannel('A')
            self.assertGreater(alpha.getextrema()[1], 170)
            self.assertEqual(alpha.getpixel((0, 0)), 0)

        for build in (_classic_docx, _new_docx):
            with self.subTest(template=build.__name__):
                docx_bytes = build(_data(approved_bytes))
                with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
                    picture_name = next(name for name in archive.namelist() if name.startswith('word/media/'))
                    with Image.open(io.BytesIO(archive.read(picture_name))) as embedded:
                        self.assertGreater(embedded.getchannel('A').getextrema()[1], 170)
                self.assertIn('15.10.2026', Document(io.BytesIO(docx_bytes)).element.xml)

    def test_low_opacity_transparent_pen_is_visible_on_white_document(self):
        canvas = Image.new('RGBA', (600, 200), (0, 0, 0, 0))
        ImageDraw.Draw(canvas).line((150, 100, 430, 95), fill=(0, 0, 0, 40), width=5)
        raw = io.BytesIO()
        canvas.save(raw, format='PNG')
        display_bytes, _ = _normalise_provider_signature_image(raw.getvalue())
        with Image.open(io.BytesIO(display_bytes)) as rendered:
            alpha = rendered.getchannel('A')
            self.assertGreater(alpha.getextrema()[1], 170)
            self.assertEqual(alpha.getpixel((0, 0)), 0)

    def test_invisible_signature_is_rejected_instead_of_generating_blank_signed_contract(self):
        raw = io.BytesIO()
        Image.new('RGBA', (100, 50), (255, 255, 255, 0)).save(raw, 'PNG')
        with self.assertRaisesRegex(ValueError, 'keine erkennbaren Linien'):
            _normalise_provider_signature_image(raw.getvalue())
