"""Single-month work-time PDF and explicit signed saldo presentation."""
import csv
import io
from datetime import date
from decimal import Decimal
from urllib.parse import unquote

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from core.models import WorkingTimeAccountRecord
from core.working_time import (
    export_csv, export_xlsx, signed_saldo, worker_pdf, worker_docx,
)


def sample_record(worker_user):
    return WorkingTimeAccountRecord.objects.create(
        worker=worker_user.worker_profile,
        year_month=date(2026, 9, 1),
        ist_hours=Decimal('8.00'),
        soll_hours=Decimal('7.00'),
        paid_total_hours=Decimal('7.00'),
        difference_hours=Decimal('1.00'),
        saldo_cumulative=Decimal('1.00'),
        raw_entries=[{
            'id': 'time-1',
            'clock_in': '2026-09-05T08:15:00+02:00',
            'clock_out': '2026-09-05T16:45:00+02:00',
            'local_clock_in': '2026-09-05T08:15:00+02:00',
            'local_clock_out': '2026-09-05T16:45:00+02:00',
            'planned_start': '2026-09-05T08:00:00+02:00',
            'planned_end': '2026-09-05T17:00:00+02:00',
            'break_minutes': 30,
            'worked_minutes': 480,
            'saturday_minutes': 120,
            'sunday_minutes': 0,
            'holiday_minutes': 0,
        }],
    )


def pdf_text(content):
    return '\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)


def test_signed_saldo_explicit_sign_for_positive_negative_and_zero():
    assert signed_saldo('1.03') == '+1,03'
    assert signed_saldo('-13.29') == '-13,29'
    assert signed_saldo('0') == '+0,00'


@pytest.mark.django_db
def test_single_month_employee_pdf_columns_and_filename(auth_admin, worker_user):
    record = sample_record(worker_user)
    result = auth_admin.get(
        f'/api/working-time/attendance-pdf/{worker_user.worker_profile.pk}/?month=2026-09'
    )
    assert result.status_code == 200
    assert result['Content-Type'] == 'application/pdf'
    assert result.content[:4] == b'%PDF'
    filename = unquote(result['Content-Disposition'].split("filename*=UTF-8''")[1])
    assert filename == 'Arbeitszeitnachweis - Anna Becker - September 2026.pdf'
    text = pdf_text(result.content)
    for word in ('Arbeitszeitnachweis', 'September 2026', 'Datum', 'Soll', 'Ist',
                 'Pause', 'Netto', 'Sa.', 'So.', 'Feiertag', 'Gesamt'):
        assert word in text
    assert '05.09.2026' in text
    assert '08:00 - 17:00' in text
    assert '08:15 - 16:45' in text
    assert '30 Min.' in text
    assert '8:00' in text
    assert '2:00' in text
    assert '0:00' in text


@pytest.mark.django_db
def test_monthly_pdf_requires_admin_and_valid_month(auth_admin, auth_worker, auth_client, worker_user):
    worker_id = worker_user.worker_profile.pk
    path = f'/api/working-time/attendance-pdf/{worker_id}/'
    assert auth_worker.get(path + '?month=2026-09').status_code == 403
    assert auth_client.get(path + '?month=2026-09').status_code == 403
    assert auth_admin.get(path + '?month=2026-13').status_code == 400
    assert auth_admin.get(path + '?month=2026-09%20junk').status_code == 400
    assert auth_admin.get(path + '?month=2026-09').status_code == 404


@pytest.mark.django_db
def test_worker_pdf_word_csv_xlsx_keep_signed_saldo_without_changing_numeric_cells(worker_user):
    record = sample_record(worker_user)
    queryset = WorkingTimeAccountRecord.objects.filter(pk=record.pk)
    pdf = pdf_text(worker_pdf(worker_user.worker_profile, queryset))
    assert '+1,00' in pdf
    assert 'Auszahlungs' in pdf
    assert 'Soll' in pdf
    doc = Document(io.BytesIO(worker_docx(worker_user.worker_profile, queryset)))
    word_cells = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert '+1,00' in word_cells
    assert 'Auszahlungsbetrag €' in word_cells
    assert 'Soll' in word_cells
    csv_content = export_csv(queryset).content.decode('utf-8-sig')
    csv_rows = list(csv.reader(io.StringIO(csv_content), delimiter=';'))
    assert csv_rows[1][8] == '+1,00'
    assert csv_rows[1][13] == '+1,00'
    workbook = load_workbook(io.BytesIO(export_xlsx(queryset).content))
    sheet = workbook['Arbeitszeitkonto']
    assert sheet['I2'].value == 1
    assert sheet['I2'].number_format == '+0.00;-0.00;+0.00'
    assert sheet['N2'].value == 1
    assert sheet['N2'].number_format == '+0.00;-0.00;+0.00'
