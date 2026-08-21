import io
from datetime import date

import fitz
import pytest
from django.core.files.base import ContentFile
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from core import document_engine
from core.models import Contract, ContractTemplate, EmployeeMasterData


def _docx_bytes():
    document = Document()
    document.add_paragraph('Arbeitsvertrag zwischen')
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _party_form_pdf():
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    pdf.setFont('Helvetica', 9)
    for label, y in [('Firma:', 760), ('Anschrift:', 725), ('Name:', 690), ('Anschrift:', 655)]:
        pdf.drawString(55, y + 2, label)
        pdf.line(125, y, 480, y)
    pdf.drawString(55, 610, '1. Vertragsgegenstand')
    pdf.save()
    return output.getvalue()


def _unfillable_pdf():
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    pdf.setFont('Helvetica', 11)
    pdf.drawString(60, 760, 'Arbeitsvertrag zwischen')
    pdf.drawString(60, 720, 'Unveränderlicher Vertragstext ohne Eingabefelder')
    pdf.save()
    return output.getvalue()


def _template(worker_user):
    fields = [
        {'name': 'company_name', 'label': 'Firma', 'source': 'company.name', 'required': True, 'type': 'text'},
        {'name': 'company_address', 'label': 'Firmenanschrift', 'source': 'company.address', 'required': True, 'type': 'text'},
        {'name': 'employee_name', 'label': 'Mitarbeiter', 'source': 'worker.full_name', 'required': True, 'type': 'text'},
        {'name': 'employee_address', 'label': 'Anschrift Mitarbeiter', 'source': 'master.full_address', 'required': True, 'type': 'text'},
    ]
    template = ContractTemplate.objects.create(
        name='Legacy SmartDocs Arbeitsvertrag',
        slug='legacy-smartdocs-arbeitsvertrag',
        kind=ContractTemplate.Kind.EMPLOYMENT,
        source_format=ContractTemplate.SourceFormat.DOCX,
        schema={'fields': fields, 'signature_roles': ['employee', 'employer']},
        requires_signature=True,
    )
    template.source_file.save('legacy-contract.docx', ContentFile(_docx_bytes()), save=True)
    EmployeeMasterData.objects.update_or_create(
        worker=worker_user.worker_profile,
        defaults={
            'data': {
                'street': 'Musterstraße 12',
                'postal_code': '60311',
                'city': 'Frankfurt am Main',
            }
        },
    )
    return template


@pytest.mark.django_db
def test_legacy_docx_blank_lines_are_populated_through_smartdocs(monkeypatch, worker_user, admin_user):
    template = _template(worker_user)
    contract = Contract.objects.create(
        template=template,
        worker=worker_user.worker_profile,
        title='Arbeitsvertrag',
        starts_on=date(2026, 9, 1),
        created_by=admin_user,
    )

    # LibreOffice conversion itself is not under test here. The regression is the
    # legacy converted PDF with printed blank lines and no {{placeholders}}.
    monkeypatch.setattr(document_engine, 'convert_docx_to_pdf', lambda _docx: _party_form_pdf())

    document_engine.generate_contract_files(contract)

    contract.refresh_from_db()
    assert contract.status == Contract.Status.READY
    assert contract.pdf
    with contract.pdf.open('rb') as handle:
        generated = handle.read()
    pdf = fitz.open(stream=generated, filetype='pdf')
    try:
        text = '\n'.join(page.get_text('text') for page in pdf)
    finally:
        pdf.close()

    assert 'A+ Solution GmbH' in text
    assert 'Carl-Sonnenschein-Str. 57, 65936 Frankfurt am Main' in text
    assert 'Anna Becker' in text
    assert 'Musterstraße 12 60311 Frankfurt am Main' in text
    assert 'smartdocs-layout-v2-docx-guard' in contract.data_snapshot['_smartdocs_pdf']


@pytest.mark.django_db
def test_unfillable_legacy_docx_is_not_marked_ready(monkeypatch, worker_user, admin_user):
    template = _template(worker_user)
    contract = Contract.objects.create(
        template=template,
        worker=worker_user.worker_profile,
        title='Arbeitsvertrag',
        starts_on=date(2026, 9, 1),
        created_by=admin_user,
    )
    monkeypatch.setattr(document_engine, 'convert_docx_to_pdf', lambda _docx: _unfillable_pdf())

    with pytest.raises(document_engine.DocumentGenerationError, match='SmartDocs'):
        document_engine.generate_contract_files(contract)

    contract.refresh_from_db()
    assert contract.status == Contract.Status.DRAFT
    assert not contract.pdf
    assert contract.generated_at is None
